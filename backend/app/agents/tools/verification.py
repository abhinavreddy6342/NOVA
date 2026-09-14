from pathlib import Path
from typing import Any, Dict, Optional
import csv

import pymupdf  # fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image


def verify_file_basic(file_path: Path) -> tuple[bool, str]:
    """Verify that a file exists and is non-empty."""
    if not file_path.exists():
        return False, f"File does not exist: {file_path.name}"
    if not file_path.is_file():
        return False, f"Path is not a file: {file_path.name}"
    if file_path.stat().st_size <= 0:
        return False, f"File is empty (0 bytes): {file_path.name}"
    return True, "File exists and is non-empty"


def verify_pdf(file_path: Path) -> tuple[bool, str]:
    """Verify that a PDF exists, is non-empty, and can be parsed."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        doc = pymupdf.open(str(file_path))
        page_count = len(doc)
        doc.close()
        if page_count <= 0:
            return False, f"PDF file has 0 pages: {file_path.name}"
        return True, f"PDF valid ({page_count} page(s))"
    except Exception as exc:
        return False, f"PDF verification failed: {exc}"


def verify_docx(file_path: Path) -> tuple[bool, str]:
    """Verify that a DOCX document exists, is non-empty, and can be opened."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        doc = Document(str(file_path))
        p_count = len(doc.paragraphs)
        t_count = len(doc.tables)
        return True, f"DOCX valid ({p_count} paragraph(s), {t_count} table(s))"
    except Exception as exc:
        return False, f"DOCX verification failed: {exc}"


def verify_xlsx(file_path: Path) -> tuple[bool, str]:
    """Verify that an XLSX workbook exists, is non-empty, and can be loaded."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        sheet_count = len(wb.sheetnames)
        wb.close()
        if sheet_count <= 0:
            return False, f"XLSX workbook has no sheets: {file_path.name}"
        return True, f"XLSX valid ({sheet_count} sheet(s))"
    except Exception as exc:
        return False, f"XLSX verification failed: {exc}"


def verify_csv(file_path: Path) -> tuple[bool, str]:
    """Verify that a CSV file exists, is non-empty, and can be parsed."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        with file_path.open("r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f)
            row_count = sum(1 for _ in reader)
        if row_count <= 0:
            return False, f"CSV file has no rows: {file_path.name}"
        return True, f"CSV valid ({row_count} row(s))"
    except Exception as exc:
        return False, f"CSV verification failed: {exc}"


def verify_pptx(file_path: Path) -> tuple[bool, str]:
    """Verify that a PPTX presentation exists, is non-empty, and can be parsed."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        prs = Presentation(str(file_path))
        slide_count = len(prs.slides)
        if slide_count <= 0:
            return False, f"PPTX presentation has 0 slides: {file_path.name}"
        return True, f"PPTX valid ({slide_count} slide(s))"
    except Exception as exc:
        return False, f"PPTX verification failed: {exc}"


def verify_image(file_path: Path) -> tuple[bool, str]:
    """Verify that an image file exists, is non-empty, and can be opened."""
    ok, msg = verify_file_basic(file_path)
    if not ok:
        return False, msg

    try:
        with Image.open(str(file_path)) as img:
            img.verify()
        return True, f"Image valid ({file_path.name})"
    except Exception as exc:
        return False, f"Image verification failed: {exc}"


def verify_artifact(file_path: Path, artifact_type: Optional[str] = None) -> tuple[bool, str]:
    """
    Universal verification dispatcher based on extension or artifact_type.
    """
    ext = file_path.suffix.lower()
    if artifact_type:
        artifact_type = artifact_type.lower()

    if ext == ".pdf" or artifact_type == "pdf":
        return verify_pdf(file_path)
    elif ext == ".docx" or artifact_type == "docx":
        return verify_docx(file_path)
    elif ext == ".xlsx" or artifact_type == "xlsx":
        return verify_xlsx(file_path)
    elif ext == ".csv" or artifact_type == "csv":
        return verify_csv(file_path)
    elif ext == ".pptx" or artifact_type == "pptx":
        return verify_pptx(file_path)
    elif ext in {".png", ".jpg", ".jpeg", ".webp"} or artifact_type == "image":
        return verify_image(file_path)
    else:
        return verify_file_basic(file_path)
