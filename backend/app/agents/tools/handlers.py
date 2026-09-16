from pathlib import Path
from typing import Any, Dict, List
from datetime import datetime
import csv
import hashlib
import os
import re
import threading

from docx import Document
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
)
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import (
    Inches,
    Pt,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openpyxl import load_workbook
import openpyxl

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

from openpyxl.utils import (
    get_column_letter,
)

from PIL import Image, ImageDraw

from app.knowledge.rag import retrieve_context
from app.services.knowledge.extractor import extract_document
from app.agents.tools.code_sandbox import (
    execute_python_in_sandbox,
)
from app.agents.tools.spreadsheet_analyzer import (
    analyze_spreadsheet,
)

from app.agents.tools.verification import (
    verify_artifact,
)

from pptx import Presentation
from pptx.util import Inches as PPTXInches, Pt as PPTXPt
from pptx.dml.color import RGBColor
from pptx.enum.text import (
    PP_ALIGN,
    MSO_ANCHOR,
)
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as RLImage,
)
from reportlab.pdfgen import canvas


# ---------------------------------------------------------------------------
# NOVA WORKSPACE CONFIGURATION
# ---------------------------------------------------------------------------

BACKEND_ROOT = Path(__file__).resolve().parents[3]

WORKSPACE_ROOT = (
    BACKEND_ROOT / "app" / "workspace"
).resolve()

ALLOWED_READ_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".css",
    ".xml",
    ".yaml",
    ".yml",
    ".log",
}

ALLOWED_WRITE_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".css",
    ".xml",
    ".yaml",
    ".yml",
    ".log",
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".png",
    ".jpg",
    ".jpeg",
}

ALLOWED_DOCUMENT_WRITE_EXTENSIONS = {
    ".docx",
}

ALLOWED_SPREADSHEET_EXTENSIONS = {
    ".xlsx",
    ".csv",
}

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_WRITE_SIZE = 5 * 1024 * 1024
MAX_DOCUMENT_SIZE = 10 * 1024 * 1024

MAX_SPREADSHEET_ROWS = 500
MAX_SPREADSHEET_COLUMNS = 50

MAX_CODE_SIZE = 100_000


# ---------------------------------------------------------------------------
# WORKSPACE PATH SECURITY
# ---------------------------------------------------------------------------

def _resolve_workspace_file(
    file_path: str,
) -> Path:
    """
    Resolve a file path and ensure it remains inside
    NOVA's controlled workspace.
    """

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "file_path is required."
        )

    raw_path = Path(
        str(file_path).strip()
    )

    if raw_path.is_absolute():
        resolved = raw_path.resolve()
    else:
        resolved = (
            WORKSPACE_ROOT / raw_path
        ).resolve()

    try:
        resolved.relative_to(
            WORKSPACE_ROOT
        )
    except ValueError as exc:
        raise PermissionError(
            "Access denied: file must remain inside NOVA's controlled workspace."
        ) from exc

    return resolved


# ---------------------------------------------------------------------------
# DOCUMENT READER
# ---------------------------------------------------------------------------

def document_reader_handler(
    file_path: str,
    **kwargs: Any,
) -> Dict[str, Any]:
    """
    Read and extract a supported local document.

    Supported:
    - PDF
    - TXT
    - DOCX
    """

    if not file_path:
        raise ValueError(
            "document_reader requires a file_path."
        )

    result = extract_document(
        file_path
    )

    return result.to_dict()


# ---------------------------------------------------------------------------
# DOCUMENT WRITER HELPERS
# ---------------------------------------------------------------------------

def _clean_document_text(
    content: str,
) -> str:
    """
    Remove accidental Markdown code fences and normalize escaped
    Markdown characters before DOCX rendering.
    """

    if content is None:
        return ""

    text = str(content).strip()

    if (
        text.startswith("```")
        and text.endswith("```")
    ):
        lines = text.splitlines()

        if lines:
            lines = lines[1:]

        if (
            lines
            and lines[-1].strip() == "```"
        ):
            lines = lines[:-1]

        text = "\n".join(
            lines
        ).strip()

    replacements = (
        (r"\*\*", "**"),
        (r"\_\_", "__"),
        (r"\*", "*"),
        (r"\_", "_"),
        (r"\`", "`"),
        (r"\#", "#"),
        (r"\-", "-"),
    )

    for old, new in replacements:
        text = text.replace(
            old,
            new,
        )

    return text


def _clean_inline_markdown(
    text: str,
) -> str:
    """
    Normalize common escaped Markdown sequences.
    """

    if text is None:
        return ""

    value = str(text)

    replacements = (
        (r"\*\*", "**"),
        (r"\_\_", "__"),
        (r"\*", "*"),
        (r"\_", "_"),
        (r"\`", "`"),
        (r"\#", "#"),
        (r"\-", "-"),
        (r"\.", "."),
        (r"\:", ":"),
        (r"\(", "("),
        (r"\)", ")"),
    )

    for old, new in replacements:
        value = value.replace(
            old,
            new,
        )

    return value.strip()


def _add_inline_runs(
    paragraph: Any,
    text: str,
) -> None:
    """
    Convert common inline Markdown syntax into DOCX runs.
    """

    normalized = _clean_inline_markdown(
        text
    )

    if not normalized:
        return

    pattern = re.compile(
        r"("
        r"\*\*(.+?)\*\*"
        r"|__(.+?)__"
        r"|(?<!\*)\*(?!\s)(.+?)(?<!\s)\*"
        r"|(?<!_)_(?!\s)(.+?)(?<!\s)_"
        r"|`(.+?)`"
        r")",
        flags=re.DOTALL,
    )

    cursor = 0

    for match in pattern.finditer(
        normalized
    ):
        start, end = match.span()

        if start > cursor:
            plain_run = paragraph.add_run(
                normalized[
                    cursor:start
                ]
            )

            plain_run.font.name = "Arial"
            plain_run.font.size = Pt(
                10.5
            )

        matched = match.group(0)

        if (
            matched.startswith("**")
            and matched.endswith("**")
        ):
            inner = matched[2:-2]

            run = paragraph.add_run(
                inner
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(
                10.5
            )

        elif (
            matched.startswith("__")
            and matched.endswith("__")
        ):
            inner = matched[2:-2]

            run = paragraph.add_run(
                inner
            )

            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(
                10.5
            )

        elif (
            matched.startswith("*")
            and matched.endswith("*")
        ):
            inner = matched[1:-1]

            run = paragraph.add_run(
                inner
            )

            run.italic = True
            run.font.name = "Arial"
            run.font.size = Pt(
                10.5
            )

        elif (
            matched.startswith("_")
            and matched.endswith("_")
        ):
            inner = matched[1:-1]

            run = paragraph.add_run(
                inner
            )

            run.italic = True
            run.font.name = "Arial"
            run.font.size = Pt(
                10.5
            )

        elif (
            matched.startswith("`")
            and matched.endswith("`")
        ):
            inner = matched[1:-1]

            run = paragraph.add_run(
                inner
            )

            run.font.name = "Consolas"
            run.font.size = Pt(
                9.5
            )

            run.font.color.rgb = (
                __import__(
                    "docx"
                ).shared.RGBColor(
                    45,
                    55,
                    72,
                )
            )

        else:
            run = paragraph.add_run(
                matched
            )

            run.font.name = "Arial"
            run.font.size = Pt(
                10.5
            )

        cursor = end

    if cursor < len(normalized):
        final_run = paragraph.add_run(
            normalized[
                cursor:
            ]
        )

        final_run.font.name = "Arial"
        final_run.font.size = Pt(
            10.5
        )


def _remove_paragraph_border(
    paragraph: Any,
) -> None:
    """
    Safely remove inherited paragraph borders.
    """

    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(
        qn("w:pBdr")
    )

    if p_bdr is not None:
        p_pr.remove(
            p_bdr
        )


def _add_horizontal_rule(
    document: Document,
) -> None:
    """
    Add a clean horizontal rule to a DOCX.
    """

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(
        4
    )

    paragraph.paragraph_format.space_after = Pt(
        10
    )

    p_pr = paragraph._p.get_or_add_pPr()

    p_bdr = OxmlElement(
        "w:pBdr"
    )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single",
    )

    bottom.set(
        qn("w:sz"),
        "6",
    )

    bottom.set(
        qn("w:space"),
        "1",
    )

    bottom.set(
        qn("w:color"),
        "0EA5E9",
    )

    p_bdr.append(
        bottom
    )

    p_pr.append(
        p_bdr
    )


def _set_cell_shading(
    cell: Any,
    fill: str,
) -> None:
    """
    Apply background shading to a DOCX table cell.
    """

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:
        shd = OxmlElement(
            "w:shd"
        )

        tc_pr.append(
            shd
        )

    shd.set(
        qn("w:fill"),
        fill,
    )


def _style_document(
    document: Document,
) -> None:
    """
    Configure professional NOVA document typography and layout.
    """

    section = document.sections[0]

    section.top_margin = Inches(
        0.75
    )

    section.bottom_margin = Inches(
        0.75
    )

    section.left_margin = Inches(
        0.85
    )

    section.right_margin = Inches(
        0.85
    )

    section.header_distance = Inches(
        0.35
    )

    section.footer_distance = Inches(
        0.35
    )

    normal_style = document.styles[
        "Normal"
    ]

    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(
        10.5
    )

    normal_style.paragraph_format.space_after = Pt(
        7
    )

    normal_style.paragraph_format.line_spacing = 1.12

    normal_style.paragraph_format.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    for style_name, size, color in (
        (
            "Title",
            22,
            "0F172A",
        ),
        (
            "Heading 1",
            16,
            "0284C7",
        ),
        (
            "Heading 2",
            13,
            "0F172A",
        ),
        (
            "Heading 3",
            11.5,
            "0284C7",
        ),
    ):
        if style_name not in document.styles:
            continue

        style = document.styles[
            style_name
        ]

        style.font.name = "Arial"
        style.font.size = Pt(
            size
        )

        style.font.color.rgb = (
            __import__(
                "docx"
            ).shared.RGBColor.from_string(
                color
            )
        )

        style.font.bold = True

    if "NOVA Body" not in document.styles:
        nova_body = document.styles.add_style(
            "NOVA Body",
            WD_STYLE_TYPE.PARAGRAPH,
        )

        nova_body.font.name = "Arial"
        nova_body.font.size = Pt(
            10.5
        )

        nova_body.paragraph_format.space_after = Pt(
            6
        )

        nova_body.paragraph_format.line_spacing = 1.12

    header = section.header

    if len(header.paragraphs) == 0:
        header_paragraph = header.add_paragraph()
    else:
        header_paragraph = header.paragraphs[0]

    header_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    header_run = header_paragraph.add_run(
        "NOVA | Sovereign Industrial Intelligence"
    )

    header_run.font.name = "Arial"
    header_run.font.size = Pt(
        8
    )

    header_run.font.bold = True

    header_run.font.color.rgb = (
        __import__(
            "docx"
        ).shared.RGBColor(
            100,
            116,
            139,
        )
    )

    footer = section.footer

    if len(footer.paragraphs) == 0:
        footer_paragraph = footer.add_paragraph()
    else:
        footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "Confidential — Generated locally by NOVA"
    )

    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(
        8
    )

    footer_run.font.color.rgb = (
        __import__(
            "docx"
        ).shared.RGBColor(
            100,
            116,
            139,
        )
    )


def _add_document_title(
    document: Document,
    title: str,
) -> None:
    """
    Add one clean professional document title.
    """

    clean_title = _clean_inline_markdown(
        title
    )

    if not clean_title:
        return

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    paragraph.paragraph_format.space_before = Pt(
        4
    )

    paragraph.paragraph_format.space_after = Pt(
        8
    )

    run = paragraph.add_run(
        clean_title
    )

    run.font.name = "Arial"
    run.font.size = Pt(
        22
    )

    run.font.bold = True

    run.font.color.rgb = (
        __import__(
            "docx"
        ).shared.RGBColor(
            15,
            23,
            42,
        )
    )

    accent = document.add_paragraph()

    accent.paragraph_format.space_after = Pt(
        14
    )

    p_pr = accent._p.get_or_add_pPr()

    p_bdr = OxmlElement(
        "w:pBdr"
    )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single",
    )

    bottom.set(
        qn("w:sz"),
        "12",
    )

    bottom.set(
        qn("w:space"),
        "1",
    )

    bottom.set(
        qn("w:color"),
        "0EA5E9",
    )

    p_bdr.append(
        bottom
    )

    p_pr.append(
        p_bdr
    )


def _extract_numbered_item(
    line: str,
) -> str | None:

    match = re.match(
        r"^\s*\d+\.\s+(.+)$",
        line,
    )

    if not match:
        return None

    return match.group(
        1
    ).strip()


def _extract_bullet_item(
    line: str,
) -> str | None:

    match = re.match(
        r"^\s*[-*+]\s+(.+)$",
        line,
    )

    if not match:
        return None

    return match.group(
        1
    ).strip()


def _clean_heading_text(
    line: str,
) -> str:

    cleaned = re.sub(
        r"^\s*#{1,6}\s+",
        "",
        line,
    )

    cleaned = _clean_inline_markdown(
        cleaned
    )

    cleaned = cleaned.strip(
        "*_` "
    )

    return cleaned


def _add_numbered_item(
    document: Document,
    text: str,
) -> None:

    paragraph = document.add_paragraph(
        style="List Number"
    )

    paragraph.paragraph_format.space_after = Pt(
        5
    )

    paragraph.paragraph_format.left_indent = Inches(
        0.25
    )

    paragraph.paragraph_format.first_line_indent = Inches(
        -0.15
    )

    _add_inline_runs(
        paragraph,
        text,
    )


def _add_bullet_item(
    document: Document,
    text: str,
) -> None:

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.paragraph_format.space_after = Pt(
        4
    )

    paragraph.paragraph_format.left_indent = Inches(
        0.25
    )

    paragraph.paragraph_format.first_line_indent = Inches(
        -0.15
    )

    _add_inline_runs(
        paragraph,
        text,
    )


def _add_normal_paragraph(
    document: Document,
    text: str,
) -> None:

    paragraph = document.add_paragraph(
        style="NOVA Body"
    )

    paragraph.paragraph_format.space_after = Pt(
        7
    )

    paragraph.paragraph_format.line_spacing = 1.12

    _add_inline_runs(
        paragraph,
        text,
    )


def _add_content_to_document(
    document: Document,
    content: str,
) -> None:

    cleaned = _clean_document_text(
        content
    )

    if not cleaned:
        document.add_paragraph(
            "No content was provided."
        )
        return

    lines = cleaned.splitlines()

    index = 0

    while index < len(lines):
        raw_line = lines[index]

        line = raw_line.strip()

        if not line:
            index += 1
            continue

        if re.match(
            r"^\s*([-*_])(?:\s*\1){2,}\s*$",
            line,
        ):
            _add_horizontal_rule(
                document
            )

            index += 1
            continue

        if re.match(
            r"^\s*###\s+",
            line,
        ):
            heading = document.add_heading(
                _clean_heading_text(
                    line
                ),
                level=3,
            )

            heading.paragraph_format.space_before = Pt(
                10
            )

            heading.paragraph_format.space_after = Pt(
                5
            )

            index += 1
            continue

        if re.match(
            r"^\s*##\s+",
            line,
        ):
            heading = document.add_heading(
                _clean_heading_text(
                    line
                ),
                level=2,
            )

            heading.paragraph_format.space_before = Pt(
                12
            )

            heading.paragraph_format.space_after = Pt(
                6
            )

            index += 1
            continue

        if re.match(
            r"^\s*#\s+",
            line,
        ):
            heading = document.add_heading(
                _clean_heading_text(
                    line
                ),
                level=1,
            )

            heading.paragraph_format.space_before = Pt(
                14
            )

            heading.paragraph_format.space_after = Pt(
                7
            )

            index += 1
            continue

        numbered_item = _extract_numbered_item(
            line
        )

        if numbered_item is not None:
            _add_numbered_item(
                document,
                numbered_item,
            )

            index += 1
            continue

        bullet_item = _extract_bullet_item(
            line
        )

        if bullet_item is not None:
            _add_bullet_item(
                document,
                bullet_item,
            )

            index += 1
            continue

        paragraph_lines = [
            line
        ]

        next_index = index + 1

        while next_index < len(lines):
            candidate = lines[
                next_index
            ].strip()

            if not candidate:
                break

            if re.match(
                r"^\s*#{1,6}\s+",
                candidate,
            ):
                break

            if _extract_numbered_item(
                candidate
            ) is not None:
                break

            if _extract_bullet_item(
                candidate
            ) is not None:
                break

            if re.match(
                r"^\s*([-*_])(?:\s*\1){2,}\s*$",
                candidate,
            ):
                break

            paragraph_lines.append(
                candidate
            )

            next_index += 1

        paragraph_text = " ".join(
            paragraph_lines
        ).strip()

        _add_normal_paragraph(
            document,
            paragraph_text,
        )

        index = next_index


def _synthesize_evidence_report_from_context(
    context: Dict[str, Any],
    title: str = "File Review Report",
) -> str:
    """
    Synthesize a structured Markdown report from completed reader step results in _context.
    """
    if not isinstance(context, dict):
        return "No evidence reader results available."

    steps = context.get("steps", {})
    if not isinstance(steps, dict) or not steps:
        return "No completed file reader steps were found."

    lines: List[str] = [
        f"# {title or 'File Review Report'}",
        "",
        "## Executive Summary",
        "This report synthesizes the detailed evidence review of all supplied workspace files. "
        "Each file was inspected and analyzed using NOVA's local reader engines based on its actual extracted content.",
        "",
        "## File-by-File Evidence Analysis",
    ]

    file_index = 1
    for step_id, result in steps.items():
        if not isinstance(result, dict):
            continue

        file_name = (
            result.get("file_name")
            or result.get("filename")
            or Path(result.get("file_path", f"file_{file_index}")).name
        )
        ext = Path(file_name).suffix.lower()

        lines.append(f"\n### {file_index}. {file_name}")
        file_index += 1

        # Check for spreadsheet structure
        if "sheets" in result and isinstance(result["sheets"], list):
            sheet_count = result.get("sheet_count", len(result["sheets"]))
            lines.append(f"- **What it is**: Excel/CSV Workbook containing {sheet_count} sheet(s).")
            lines.append("- **Sheet Details**:")
            for s in result["sheets"]:
                if not isinstance(s, dict):
                    continue
                sname = s.get("sheet_name", "Sheet")
                headers = s.get("headers", [])
                rcount = s.get("row_count", 0)
                ccount = s.get("column_count", len(headers))
                lines.append(f"  - **{sname}**: {rcount} data rows x {ccount} columns.")
                if headers:
                    lines.append(f"    - Columns: {', '.join(str(h) for h in headers)}")
                rows = s.get("rows", [])
                if rows:
                    lines.append("    - Sample records:")
                    for r in rows[:4]:
                        if isinstance(r, (list, tuple)):
                            r_str = " | ".join("" if v is None else str(v) for v in r)
                        else:
                            r_str = str(r)
                        lines.append(f"      - {r_str}")
            lines.append("")
            continue

        # Extract text content for documents / text / code / images
        extracted_text = (
            result.get("content")
            or result.get("text")
            or result.get("extracted_text")
            or ""
        )
        if not isinstance(extracted_text, str):
            extracted_text = str(extracted_text)

        extracted_text = extracted_text.strip()

        if ext in {".css", ".json", ".py", ".js", ".jsx", ".ts", ".tsx", ".html", ".xml", ".yaml", ".yml", ".log"}:
            raw_lines = [l for l in extracted_text.splitlines() if l.strip()]
            line_count = len(raw_lines)
            if ext == ".css":
                selectors = re.findall(r"([\.#a-zA-Z0-9_\-\s,]+)\s*\{", extracted_text)
                clean_sel = [s.strip() for s in selectors if s.strip() and not s.strip().startswith("@")][:12]
                lines.append(f"- **What it is**: Frontend CSS Stylesheet ({line_count} lines).")
                lines.append("- **What it contains**: Defines UI theme, colors, typography, layout rules, animations, and responsive components.")
                if clean_sel:
                    lines.append(f"- **Key Selectors**: {', '.join(clean_sel)}")
            elif ext == ".json":
                try:
                    parsed = json.loads(extracted_text)
                    if isinstance(parsed, dict):
                        keys = list(parsed.keys())
                        lines.append(f"- **What it is**: JSON Data/Configuration file ({line_count} lines).")
                        lines.append(f"- **Key Fields**: {', '.join(keys[:15])}")
                    else:
                        lines.append(f"- **What it is**: JSON Data structure ({line_count} lines).")
                except Exception:
                    lines.append(f"- **What it is**: JSON file ({line_count} lines).")
            elif ext in {".py", ".js", ".ts", ".jsx", ".tsx"}:
                funcs = re.findall(r"(?:def|function|const|class)\s+([a-zA-Z0-9_]+)", extracted_text)
                func_list = sorted(list(set(funcs)))[:12]
                lines.append(f"- **What it is**: {ext.lstrip('.').upper()} Source Code module ({line_count} lines).")
                lines.append("- **What it contains**: Implementation logic for application functionality.")
                if func_list:
                    lines.append(f"- **Declared Symbols**: {', '.join(func_list)}")
            else:
                lines.append(f"- **What it is**: {ext.lstrip('.').upper()} text file ({line_count} lines).")
                lines.append(f"- **Summary**: {extracted_text[:400]}")
        else:
            doc_type = "PDF document" if ext == ".pdf" else ("Word document" if ext == ".docx" else ("Image file" if ext in {".png", ".jpg", ".jpeg"} else "Text document"))
            lines.append(f"- **What it is**: {doc_type}.")
            if extracted_text:
                preview = extracted_text[:1200].replace("\n\n", "\n")
                lines.append(f"- **What it contains**:\n{preview}")
            else:
                lines.append("- **What it contains**: No readable text extracted.")

        lines.append("")

    lines.append("## Conclusion")
    lines.append("All attached evidence files were thoroughly analyzed using NOVA's sovereign local engines. "
                 "The findings above represent the verified content extracted from each file.")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# DOCUMENT WRITER
# ---------------------------------------------------------------------------

def document_writer_handler(
    file_path: str,
    title: str = "",
    content: str = "",
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "document_writer requires a file_path."
        )

    _context = kwargs.get("_context")
    if (content == "__EVIDENCE_REVIEW_REPORT__" or not str(content or "").strip()) and isinstance(_context, dict):
        content = _synthesize_evidence_report_from_context(_context, title=title or "File Review Report")

    if content is None:
        raise ValueError(
            "document_writer requires content."
        )

    content = str(content)

    content_size = len(
        content.encode(
            "utf-8"
        )
    )

    if content_size > MAX_DOCUMENT_SIZE:
        raise ValueError(
            "Document content is too large. "
            "Maximum allowed size is 10 MB."
        )

    path = _resolve_workspace_file(
        file_path
    )

    extension = path.suffix.lower()

    if extension not in ALLOWED_DOCUMENT_WRITE_EXTENSIONS:
        raise ValueError(
            "Unsupported document type: "
            f"{extension or '[no extension]'}."
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):
        raise ValueError(
            "document_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():
        raise ValueError(
            "document_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document()

    _style_document(
        document
    )

    clean_title = (
        _clean_inline_markdown(
            str(title).strip()
        )
        if title is not None
        else ""
    )

    cleaned_content = _clean_document_text(
        content
    )

    if clean_title:
        _add_document_title(
            document,
            clean_title,
        )

    _add_content_to_document(
        document=document,
        content=cleaned_content,
    )

    document.save(
        path
    )

    if not path.exists():
        raise RuntimeError(
            "DOCX document was not created."
        )

    file_size = path.stat().st_size

    if file_size <= 0:
        raise RuntimeError(
            "Generated DOCX document is empty."
        )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": extension,
        "size_bytes": file_size,
        "created": True,
        "title": clean_title,
        "workspace": "NOVA",
        "tool": "document_writer",
    }


# ---------------------------------------------------------------------------
# KNOWLEDGE SEARCH
# ---------------------------------------------------------------------------

def knowledge_search_handler(
    query: str,
    top_k: int = 5,
    **kwargs: Any,
) -> Dict[str, Any]:

    if not query or not query.strip():
        raise ValueError(
            "knowledge_search requires a non-empty query."
        )

    top_k = int(top_k)

    if top_k < 1 or top_k > 10:
        raise ValueError(
            "top_k must be between 1 and 10."
        )

    results = retrieve_context(
        query=query.strip(),
        top_k=top_k,
    )

    return {
        "query": query.strip(),
        "results": results,
        "count": len(results),
    }


# ---------------------------------------------------------------------------
# FILE READER
# ---------------------------------------------------------------------------

def file_reader_handler(
    file_path: str,
    **kwargs: Any,
) -> Dict[str, Any]:

    path = _resolve_workspace_file(
        file_path
    )

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if not path.is_file():
        raise ValueError(
            "The specified path is not a file."
        )

    extension = path.suffix.lower()

    if extension not in ALLOWED_READ_EXTENSIONS:
        raise ValueError(
            "Unsupported file type: "
            f"{extension or '[no extension]'}"
        )

    file_size = path.stat().st_size

    if file_size > MAX_FILE_SIZE:
        raise ValueError(
            "File is too large for file_reader. "
            "Maximum allowed size is 5 MB."
        )

    try:
        content = path.read_text(
            encoding="utf-8-sig"
        )

    except UnicodeDecodeError as exc:
        raise ValueError(
            "File is not valid UTF-8 text."
        ) from exc

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ),
        "file_name": path.name,
        "extension": extension,
        "size_bytes": file_size,
        "content": content,
        "lines": content.count("\n") + (
            1 if content else 0
        ),
        "workspace": "NOVA",
    }


# ---------------------------------------------------------------------------
# FILE WRITER
# ---------------------------------------------------------------------------

def file_writer_handler(
    file_path: str,
    content: str,
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "file_writer requires a file_path."
        )

    if content is None:
        raise ValueError(
            "file_writer requires content."
        )

    content = str(content)

    path = _resolve_workspace_file(
        file_path
    )

    extension = path.suffix.lower()

    if extension not in ALLOWED_WRITE_EXTENSIONS:
        raise ValueError(
            "Unsupported write file type: "
            f"{extension or '[no extension]'}"
        )

    content_size = len(
        content.encode(
            "utf-8"
        )
    )

    if content_size > MAX_WRITE_SIZE:
        raise ValueError(
            "Content is too large for file_writer. "
            "Maximum allowed size is 5 MB."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        content,
        encoding="utf-8",
    )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ),
        "file_name": path.name,
        "extension": extension,
        "size_bytes": path.stat().st_size,
        "created": True,
        "workspace": "NOVA",
    }


# ---------------------------------------------------------------------------
# SPREADSHEET HELPERS
# ---------------------------------------------------------------------------

def _normalize_cell_value(
    value: Any,
) -> Any:

    if value is None:
        return None

    if isinstance(
        value,
        (
            str,
            int,
            float,
            bool,
        ),
    ):
        return value

    return str(value)


def _normalize_header(
    value: Any,
    index: int,
) -> str:

    if value is None:
        return f"Column {index + 1}"

    header = str(
        value
    ).strip()

    if not header:
        return f"Column {index + 1}"

    return header


def _normalize_headers(
    headers: List[Any],
) -> List[str]:

    return [
        _normalize_header(
            value,
            index,
        )
        for index, value in enumerate(
            headers[:MAX_SPREADSHEET_COLUMNS]
        )
    ]


def _normalize_row_to_width(
    row: List[Any],
    width: int,
) -> List[Any]:

    normalized = [
        _normalize_cell_value(
            value
        )
        for value in row[:width]
    ]

    if len(normalized) < width:
        normalized.extend(
            [None]
            * (
                width - len(normalized)
            )
        )

    return normalized


def _trim_rows(
    rows: List[List[Any]],
    width: int | None = None,
) -> List[List[Any]]:

    if width is None:
        width = MAX_SPREADSHEET_COLUMNS

    width = max(
        0,
        min(
            int(width),
            MAX_SPREADSHEET_COLUMNS,
        ),
    )

    trimmed: List[List[Any]] = []

    for row in rows[:MAX_SPREADSHEET_ROWS]:

        if not isinstance(
            row,
            (
                list,
                tuple,
            ),
        ):
            row = [
                row
            ]

        trimmed.append(
            _normalize_row_to_width(
                list(row),
                width,
            )
        )

    return trimmed


def _read_csv(
    path: Path,
) -> Dict[str, Any]:

    file_size = path.stat().st_size

    if file_size > MAX_FILE_SIZE:
        raise ValueError(
            "CSV file is too large. "
            "Maximum allowed size is 5 MB."
        )

    rows: List[List[Any]] = []

    try:
        with path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as csv_file:

            sample = csv_file.read(
                4096
            )

            csv_file.seek(
                0
            )

            try:
                dialect = csv.Sniffer().sniff(
                    sample,
                    delimiters=",;\t|",
                )

            except csv.Error:
                dialect = csv.excel

            reader = csv.reader(
                csv_file,
                dialect,
            )

            for row in reader:
                rows.append(
                    row
                )

                if len(rows) >= (
                    MAX_SPREADSHEET_ROWS + 1
                ):
                    break

    except UnicodeDecodeError as exc:
        raise ValueError(
            "CSV file is not valid UTF-8 text."
        ) from exc

    except csv.Error as exc:
        raise ValueError(
            "CSV file could not be parsed."
        ) from exc

    headers = (
        rows[0]
        if rows
        else []
    )

    normalized_headers = _normalize_headers(
        headers
    )

    data_rows = (
        rows[1:]
        if len(rows) > 1
        else []
    )

    normalized_rows = _trim_rows(
        data_rows,
        width=len(
            normalized_headers
        ),
    )

    return {
        "sheet_count": 1,
        "sheets": [
            {
                "sheet_name": path.stem,
                "headers": normalized_headers,
                "rows": normalized_rows,
                "row_count": len(
                    normalized_rows
                ),
                "column_count": len(
                    normalized_headers
                ),
            }
        ],
    }


def _read_xlsx(
    path: Path,
) -> Dict[str, Any]:

    file_size = path.stat().st_size

    if file_size > MAX_FILE_SIZE:
        raise ValueError(
            "XLSX file is too large. "
            "Maximum allowed size is 5 MB."
        )

    workbook = None

    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        sheets: List[Dict[str, Any]] = []

        for worksheet in workbook.worksheets:

            rows: List[List[Any]] = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                rows.append(
                    list(row)
                )

                if len(rows) >= (
                    MAX_SPREADSHEET_ROWS + 1
                ):
                    break

            headers = (
                rows[0]
                if rows
                else []
            )

            normalized_headers = _normalize_headers(
                list(headers)
            )

            data_rows = (
                rows[1:]
                if len(rows) > 1
                else []
            )

            normalized_rows = _trim_rows(
                data_rows,
                width=len(
                    normalized_headers
                ),
            )

            sheets.append(
                {
                    "sheet_name": worksheet.title,
                    "headers": normalized_headers,
                    "rows": normalized_rows,
                    "row_count": len(
                        normalized_rows
                    ),
                    "column_count": len(
                        normalized_headers
                    ),
                }
            )

        return {
            "sheet_count": len(
                sheets
            ),
            "sheets": sheets,
        }

    finally:
        if workbook is not None:
            workbook.close()


# ---------------------------------------------------------------------------
# SPREADSHEET READER
# ---------------------------------------------------------------------------

def spreadsheet_reader_handler(
    file_path: str,
    **kwargs: Any,
) -> Dict[str, Any]:

    path = _resolve_workspace_file(
        file_path
    )

    if not path.exists():
        raise FileNotFoundError(
            f"Spreadsheet not found: {file_path}"
        )

    if not path.is_file():
        raise ValueError(
            "The specified path is not a file."
        )

    extension = path.suffix.lower()

    if extension not in ALLOWED_SPREADSHEET_EXTENSIONS:
        raise ValueError(
            "Unsupported spreadsheet type: "
            f"{extension or '[no extension]'}. "
            "Supported types are CSV and XLSX."
        )

    if extension == ".csv":

        spreadsheet = _read_csv(
            path
        )

        return {
            "file_path": str(
                path.relative_to(
                    WORKSPACE_ROOT
                )
            ),
            "file_name": path.name,
            "file_type": "csv",
            "workspace": "NOVA",
            **spreadsheet,
        }

    workbook_data = _read_xlsx(
        path
    )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ),
        "file_name": path.name,
        "file_type": "xlsx",
        "workspace": "NOVA",
        **workbook_data,
    }


# ---------------------------------------------------------------------------
# SPREADSHEET ANALYSIS
# ---------------------------------------------------------------------------

def spreadsheet_analysis_handler(
    file_path: str,
    **kwargs: Any,
) -> Dict[str, Any]:

    spreadsheet = spreadsheet_reader_handler(
        file_path=file_path,
        **kwargs,
    )

    analysis = analyze_spreadsheet(
        spreadsheet
    )

    return {
        **analysis,
        "tool": "spreadsheet_analysis",
        "workspace": "NOVA",
    }


# ---------------------------------------------------------------------------
# CODE EXECUTOR
# ---------------------------------------------------------------------------

def code_executor_handler(
    language: str,
    code: str,
    timeout: int | float | None = None,
    **kwargs: Any,
) -> Dict[str, Any]:

    if not language or not str(language).strip():
        raise ValueError(
            "code_executor requires a language."
        )

    if code is None or not str(code).strip():
        raise ValueError(
            "code_executor requires code."
        )

    normalized_language = (
        str(language)
        .strip()
        .lower()
    )

    if normalized_language in {
        "py",
        "python3",
    }:
        normalized_language = "python"

    if normalized_language != "python":
        raise ValueError(
            "Unsupported language: "
            f"{language}. Currently only Python is enabled."
        )

    result = execute_python_in_sandbox(
        code=str(code),
        timeout=timeout,
    )

    return {
        **result,
        "tool": "code_executor",
        "workspace": "NOVA sandbox",
    }


# ---------------------------------------------------------------------------
# PDF WRITER
# ---------------------------------------------------------------------------

class _NumberedCanvas(canvas.Canvas):
    """
    Canvas wrapper to draw dynamic footer page numbers.
    """

    def __init__(
        self,
        *args,
        **kwargs,
    ):
        super().__init__(
            *args,
            **kwargs,
        )

        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(
            dict(
                self.__dict__
            )
        )

        self._startPage()

    def save(self):
        num_pages = len(
            self._saved_page_states
        )

        for state in self._saved_page_states:
            self.__dict__.update(
                state
            )

            self.draw_page_decorations(
                num_pages
            )

            super().showPage()

        super().save()

    def draw_page_decorations(
        self,
        page_count,
    ):
        self.saveState()

        self.setFont(
            "Helvetica",
            8,
        )

        self.setFillColor(
            colors.HexColor(
                "#718096"
            )
        )

        self.drawString(
            54,
            750,
            "NOVA Sovereign Intelligence | Workspace Output",
        )

        self.setStrokeColor(
            colors.HexColor(
                "#2D3748"
            )
        )

        self.setLineWidth(
            0.5
        )

        self.line(
            54,
            742,
            612 - 54,
            742,
        )

        self.drawString(
            54,
            36,
            "Confidential — Local On-Premise Document",
        )

        self.drawRightString(
            612 - 54,
            36,
            f"Page {self._pageNumber} of {page_count}",
        )

        self.line(
            54,
            48,
            612 - 54,
            48,
        )

        self.restoreState()


def pdf_writer_handler(
    file_path: str,
    title: str = "",
    content: str = "",
    charts: List[str] = None,
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "pdf_writer requires a file_path."
        )

    _context = kwargs.get("_context")
    if (content == "__EVIDENCE_REVIEW_REPORT__" or not str(content or "").strip()) and isinstance(_context, dict):
        content = _synthesize_evidence_report_from_context(_context, title=title or "File Review Report")

    path = _resolve_workspace_file(
        file_path
    )

    if path.suffix.lower() != ".pdf":
        path = path.with_suffix(
            ".pdf"
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):
        raise ValueError(
            "pdf_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():
        raise ValueError(
            "pdf_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=72,
        bottomMargin=72,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor(
            "#0EA5E9"
        ),
        spaceAfter=15,
    )

    h2_style = ParagraphStyle(
        "SectionH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor(
            "#0284C7"
        ),
        spaceBefore=12,
        spaceAfter=8,
    )

    body_style = ParagraphStyle(
        "BodyDark",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor(
            "#1E293B"
        ),
        spaceAfter=8,
    )

    bullet_style = ParagraphStyle(
        "BulletDark",
        parent=body_style,
        leftIndent=15,
        spaceAfter=4,
    )

    story = []

    clean_title = (
        str(title).strip()
        if title
        else path.stem.replace(
            "_",
            " ",
        ).title()
    )

    story.append(
        Paragraph(
            clean_title,
            title_style,
        )
    )

    story.append(
        Spacer(
            1,
            10,
        )
    )

    raw_content = _clean_document_text(
        content
    )

    if raw_content:

        for line in raw_content.splitlines():

            line_str = line.strip()

            if not line_str:
                story.append(
                    Spacer(
                        1,
                        6,
                    )
                )

                continue

            if (
                line_str.startswith(
                    "### "
                )
                or line_str.startswith(
                    "## "
                )
            ):

                story.append(
                    Paragraph(
                        _clean_heading_text(
                            line_str
                        ),
                        h2_style,
                    )
                )

            elif (
                line_str.startswith("- ")
                or line_str.startswith("* ")
            ):

                story.append(
                    Paragraph(
                        (
                            f"• "
                            f"{_clean_inline_markdown(line_str[2:].strip())}"
                        ),
                        bullet_style,
                    )
                )

            else:

                story.append(
                    Paragraph(
                        _clean_inline_markdown(
                            line_str
                        ),
                        body_style,
                    )
                )

    if charts:

        for chart_rel in charts:

            try:
                chart_p = _resolve_workspace_file(
                    chart_rel
                )

                if chart_p.exists():

                    story.append(
                        Spacer(
                            1,
                            12,
                        )
                    )

                    story.append(
                        RLImage(
                            str(chart_p),
                            width=450,
                            height=250,
                        )
                    )

            except Exception:
                pass

    doc.build(
        story,
        canvasmaker=_NumberedCanvas,
    )

    verified, msg = verify_artifact(
        path,
        "pdf",
    )

    if not verified:
        raise RuntimeError(
            f"PDF creation failed verification: {msg}"
        )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": ".pdf",
        "size_bytes": path.stat().st_size,
        "created": True,
        "verification": msg,
        "workspace": "NOVA",
        "tool": "pdf_writer",
    }


# ---------------------------------------------------------------------------
# SPREADSHEET REPORT HELPERS
# ---------------------------------------------------------------------------

def _safe_sheet_name(
    name: str,
) -> str:

    cleaned = str(
        name or "Sheet"
    )

    for character in (
        "\\",
        "/",
        "*",
        "?",
        ":",
        "[",
        "]",
    ):
        cleaned = cleaned.replace(
            character,
            " ",
        )

    cleaned = " ".join(
        cleaned.split()
    ).strip()

    if not cleaned:
        cleaned = "Sheet"

    return cleaned[:31]


def _numeric_value(
    value: Any,
) -> Any:

    if value is None:
        return None

    try:
        if isinstance(
            value,
            float,
        ):
            if value != value:
                return None

            if value in (
                float("inf"),
                float("-inf"),
            ):
                return None

        return value

    except Exception:
        return value


def _build_spreadsheet_report_sheets(
    source_data: Dict[str, Any],
    analysis_data: Dict[str, Any],
    title: str,
) -> List[Dict[str, Any]]:

    sheets: List[Dict[str, Any]] = []

    analysis_summary = analysis_data.get(
        "analysis",
        {},
    )

    source_file_name = (
        source_data.get(
            "file_name"
        )
        or analysis_data.get(
            "file_name"
        )
        or "Unknown spreadsheet"
    )

    source_file_path = (
        source_data.get(
            "file_path"
        )
        or analysis_data.get(
            "file_path"
        )
        or ""
    )

    domain = (
        analysis_summary.get(
            "domain_characterization"
        )
        or "Structured Dataset"
    )

    total_rows = analysis_summary.get(
        "total_rows_analyzed",
        0,
    )

    total_missing = analysis_summary.get(
        "total_missing_values",
        0,
    )

    sheet_count = analysis_data.get(
        "sheet_count",
        len(
            source_data.get(
                "sheets",
                [],
            )
        ),
    )

    summary_rows = [
        [
            "Report Title",
            title or "NOVA Spreadsheet Analysis Report",
        ],
        [
            "Source File",
            source_file_name,
        ],
        [
            "Source Workspace Path",
            source_file_path,
        ],
        [
            "Workbook Sheets",
            sheet_count,
        ],
        [
            "Rows Analyzed",
            total_rows,
        ],
        [
            "Missing Values",
            total_missing,
        ],
        [
            "Detected Domain",
            domain,
        ],
        [
            "Numeric Columns Detected",
            len(
                analysis_summary.get(
                    "numeric_columns_detected",
                    [],
                )
            ),
        ],
        [
            "Text Columns Detected",
            len(
                analysis_summary.get(
                    "text_columns_detected",
                    [],
                )
            ),
        ],
    ]

    highest = analysis_summary.get(
        "highest_numeric_value"
    )

    lowest = analysis_summary.get(
        "lowest_numeric_value"
    )

    if highest:

        summary_rows.append(
            [
                "Highest Numeric Value",
                (
                    f"{highest.get('value')} "
                    f"({highest.get('column')})"
                ),
            ]
        )

    if lowest:

        summary_rows.append(
            [
                "Lowest Numeric Value",
                (
                    f"{lowest.get('value')} "
                    f"({lowest.get('column')})"
                ),
            ]
        )

    sheets.append(
        {
            "sheet_name": "Executive Summary",
            "headers": [
                "Metric",
                "Value",
            ],
            "rows": summary_rows,
        }
    )

    finding_rows: List[List[Any]] = []

    if highest:

        finding_rows.append(
            [
                "Highest numeric value",
                highest.get(
                    "value"
                ),
                highest.get(
                    "column"
                ),
            ]
        )

    if lowest:

        finding_rows.append(
            [
                "Lowest numeric value",
                lowest.get(
                    "value"
                ),
                lowest.get(
                    "column"
                ),
            ]
        )

    grouped_summaries = []

    for analyzed_sheet in analysis_data.get(
        "sheets",
        [],
    ):

        sheet_name = analyzed_sheet.get(
            "sheet_name",
            "Unknown",
        )

        insights = analyzed_sheet.get(
            "insights",
            {},
        )

        missing = insights.get(
            "total_missing_values",
            0,
        )

        duplicates = insights.get(
            "duplicate_rows",
            0,
        )

        if missing:

            finding_rows.append(
                [
                    "Missing values detected",
                    missing,
                    sheet_name,
                ]
            )

        if duplicates:

            finding_rows.append(
                [
                    "Duplicate rows detected",
                    duplicates,
                    sheet_name,
                ]
            )

        grouped = insights.get(
            "grouped_summaries",
            {},
        )

        if grouped:

            for group_name, records in grouped.items():

                grouped_summaries.append(
                    (
                        sheet_name,
                        group_name,
                        records,
                    )
                )

    if not finding_rows:

        finding_rows.append(
            [
                "Analysis status",
                "Completed successfully",
                "NOVA",
            ]
        )

    sheets.append(
        {
            "sheet_name": "Key Findings",
            "headers": [
                "Finding",
                "Value",
                "Source",
            ],
            "rows": finding_rows,
        }
    )

    profile_rows: List[List[Any]] = []

    for analyzed_sheet in analysis_data.get(
        "sheets",
        [],
    ):

        insights = analyzed_sheet.get(
            "insights",
            {},
        )

        profile_rows.append(
            [
                analyzed_sheet.get(
                    "sheet_name",
                    "Unknown",
                ),
                analyzed_sheet.get(
                    "row_count",
                    0,
                ),
                analyzed_sheet.get(
                    "column_count",
                    0,
                ),
                ", ".join(
                    map(
                        str,
                        insights.get(
                            "numeric_columns",
                            [],
                        ),
                    )
                ),
                ", ".join(
                    map(
                        str,
                        insights.get(
                            "text_columns",
                            [],
                        ),
                    )
                ),
                ", ".join(
                    map(
                        str,
                        insights.get(
                            "date_columns",
                            [],
                        ),
                    )
                ),
                insights.get(
                    "total_missing_values",
                    0,
                ),
                insights.get(
                    "duplicate_rows",
                    0,
                ),
                insights.get(
                    "domain_characterization",
                    "",
                ),
            ]
        )

    sheets.append(
        {
            "sheet_name": "Sheet Profiles",
            "headers": [
                "Sheet",
                "Rows",
                "Columns",
                "Numeric Columns",
                "Text Columns",
                "Date Columns",
                "Missing Values",
                "Duplicate Rows",
                "Domain",
            ],
            "rows": profile_rows,
        }
    )

    numeric_rows: List[List[Any]] = []

    for analyzed_sheet in analysis_data.get(
        "sheets",
        [],
    ):

        sheet_name = analyzed_sheet.get(
            "sheet_name",
            "Unknown",
        )

        numeric_summary = analyzed_sheet.get(
            "numeric_summary",
            {},
        )

        for column_name, stats in numeric_summary.items():

            numeric_rows.append(
                [
                    sheet_name,
                    column_name,
                    stats.get(
                        "count"
                    ),
                    _numeric_value(
                        stats.get(
                            "sum"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "average"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "median"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "minimum"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "maximum"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "range"
                        )
                    ),
                    _numeric_value(
                        stats.get(
                            "std_dev"
                        )
                    ),
                    ", ".join(
                        str(value)
                        for value in stats.get(
                            "top_values",
                            [],
                        )
                    ),
                ]
            )

    if not numeric_rows:

        numeric_rows.append(
            [
                "N/A",
                "No numeric columns detected",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "",
            ]
        )

    sheets.append(
        {
            "sheet_name": "Numeric Analysis",
            "headers": [
                "Sheet",
                "Column",
                "Count",
                "Sum",
                "Average",
                "Median",
                "Minimum",
                "Maximum",
                "Range",
                "Std Dev",
                "Top Values",
            ],
            "rows": numeric_rows,
        }
    )

    grouped_rows: List[List[Any]] = []

    for (
        sheet_name,
        group_name,
        records,
    ) in grouped_summaries:

        for record in records:

            keys = [
                key
                for key in record.keys()
                if key not in {
                    "count",
                    "sum",
                    "mean",
                }
            ]

            group_value = (
                record.get(
                    keys[0]
                )
                if keys
                else ""
            )

            grouped_rows.append(
                [
                    sheet_name,
                    group_name,
                    group_value,
                    record.get(
                        "count"
                    ),
                    _numeric_value(
                        record.get(
                            "sum"
                        )
                    ),
                    _numeric_value(
                        record.get(
                            "mean"
                        )
                    ),
                ]
            )

    if not grouped_rows:

        grouped_rows.append(
            [
                "N/A",
                "No grouped summary available",
                "",
                None,
                None,
                None,
            ]
        )

    sheets.append(
        {
            "sheet_name": "Grouped Summary",
            "headers": [
                "Sheet",
                "Grouping",
                "Group Value",
                "Count",
                "Sum",
                "Average",
            ],
            "rows": grouped_rows,
        }
    )

    for source_sheet in source_data.get(
        "sheets",
        [],
    ):

        source_sheet_name = _safe_sheet_name(
            (
                f"Source - "
                f"{source_sheet.get('sheet_name', 'Data')}"
            )
        )

        sheets.append(
            {
                "sheet_name": source_sheet_name,
                "headers": source_sheet.get(
                    "headers",
                    [],
                ),
                "rows": source_sheet.get(
                    "rows",
                    [],
                ),
            }
        )

    return sheets


# ---------------------------------------------------------------------------
# SPREADSHEET WRITER
# ---------------------------------------------------------------------------

def _find_attached_spreadsheet_from_context(
    context: Any,
) -> str | None:

    if not isinstance(
        context,
        dict,
    ):
        return None

    plan_objective = str(
        context.get(
            "plan_objective",
            "",
        )
    )

    if not plan_objective:
        return None

    matches = re.findall(
        r"(?<![A-Za-z0-9_./\\-])"
        r"(input[\\/][A-Za-z0-9_.\-\\/]+\.(?:xlsx|csv))",
        plan_objective,
        flags=re.IGNORECASE,
    )

    if not matches:
        return None

    candidate = (
        matches[-1]
        .strip()
        .rstrip(
            ".,;:)"
        )
        .replace(
            "\\",
            "/",
        )
    )

    return candidate


def spreadsheet_writer_handler(
    file_path: str,
    title: str = "",
    sheets_data: List[Dict[str, Any]] = None,
    content: str = "",
    source_file_path: str = "",
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "spreadsheet_writer requires a file_path."
        )

    path = _resolve_workspace_file(
        file_path
    )

    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(
            ".xlsx"
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):
        raise ValueError(
            "spreadsheet_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():
        raise ValueError(
            "spreadsheet_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    execution_context = kwargs.get(
        "_context",
        {},
    )

    if not source_file_path:

        source_file_path = (
            _find_attached_spreadsheet_from_context(
                execution_context
            )
            or ""
        )

    source_file_path = str(
        source_file_path or ""
    ).strip().replace(
        "\\",
        "/",
    )

    report_sheets = (
        sheets_data
        if isinstance(
            sheets_data,
            list,
        ) and sheets_data
        else None
    )

    if (
        not report_sheets
        and source_file_path
    ):

        source_extension = Path(
            source_file_path
        ).suffix.lower()

        if source_extension not in {
            ".xlsx",
            ".csv",
        }:
            raise ValueError(
                "spreadsheet_writer source_file_path must reference "
                "a CSV or XLSX file."
            )

        source_path = _resolve_workspace_file(
            source_file_path
        )

        if not source_path.exists():
            raise FileNotFoundError(
                f"Source spreadsheet not found: {source_file_path}"
            )

        source_data = spreadsheet_reader_handler(
            file_path=source_file_path
        )

        analysis_data = spreadsheet_analysis_handler(
            file_path=source_file_path
        )

        report_sheets = _build_spreadsheet_report_sheets(
            source_data=source_data,
            analysis_data=analysis_data,
            title=(
                str(title).strip()
                if title
                else "NOVA Spreadsheet Analysis Report"
            ),
        )

    if not report_sheets:

        report_sheets = [
            {
                "sheet_name": "Summary",
                "headers": [
                    "Item",
                    "Value",
                ],
                "rows": [
                    [
                        "Status",
                        "Complete",
                    ],
                    [
                        "Title",
                        title or path.stem,
                    ],
                    [
                        "Note",
                        (
                            "No source spreadsheet "
                            "was supplied for analysis."
                        ),
                    ],
                ],
            }
        ]

    wb = openpyxl.Workbook()

    wb.remove(
        wb.active
    )

    wb.properties.title = (
        str(title).strip()
        if title
        else "NOVA Spreadsheet Analysis Report"
    )

    wb.properties.subject = (
        "Generated locally by NOVA"
    )

    wb.properties.creator = (
        "NOVA Sovereign Intelligence"
    )

    title_fill = PatternFill(
        start_color="0F172A",
        end_color="0F172A",
        fill_type="solid",
    )

    header_fill = PatternFill(
        start_color="1E293B",
        end_color="1E293B",
        fill_type="solid",
    )

    accent_fill = PatternFill(
        start_color="0EA5E9",
        end_color="0EA5E9",
        fill_type="solid",
    )

    title_font = Font(
        name="Calibri",
        size=16,
        bold=True,
        color="FFFFFF",
    )

    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="00F0FF",
    )

    body_font = Font(
        name="Calibri",
        size=10,
        color="0F172A",
    )

    accent_font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color="FFFFFF",
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color="CBD5E1",
        ),
        right=Side(
            style="thin",
            color="CBD5E1",
        ),
        top=Side(
            style="thin",
            color="CBD5E1",
        ),
        bottom=Side(
            style="thin",
            color="CBD5E1",
        ),
    )

    used_names = set()

    for index, sheet_info in enumerate(
        report_sheets
    ):

        raw_name = str(
            sheet_info.get(
                "sheet_name",
                f"Sheet {index + 1}",
            )
        )

        sheet_name = _safe_sheet_name(
            raw_name
        )

        base_name = sheet_name
        counter = 2

        while sheet_name.lower() in used_names:

            suffix = f" {counter}"

            sheet_name = (
                base_name[
                    :(
                        31
                        - len(suffix)
                    )
                ]
                + suffix
            )

            counter += 1

        used_names.add(
            sheet_name.lower()
        )

        ws = wb.create_sheet(
            title=sheet_name
        )

        headers = sheet_info.get(
            "headers",
            [],
        )

        rows = sheet_info.get(
            "rows",
            [],
        )

        headers = [
            _normalize_cell_value(
                value
            )
            for value in headers
        ]

        if headers:

            if sheet_name == "Executive Summary":

                title_text = (
                    str(title).strip()
                    if title
                    else "NOVA Spreadsheet Analysis Report"
                )

                last_col = max(
                    2,
                    len(headers),
                )

                ws.merge_cells(
                    start_row=1,
                    start_column=1,
                    end_row=1,
                    end_column=last_col,
                )

                title_cell = ws.cell(
                    row=1,
                    column=1,
                    value=title_text,
                )

                title_cell.fill = (
                    title_fill
                )

                title_cell.font = (
                    title_font
                )

                title_cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                )

                ws.row_dimensions[
                    1
                ].height = 28

                header_row = 3

            else:
                header_row = 1

            for col_idx, header in enumerate(
                headers,
                start=1,
            ):

                cell = ws.cell(
                    row=header_row,
                    column=col_idx,
                    value=header,
                )

                cell.fill = (
                    accent_fill
                    if sheet_name == "Key Findings"
                    else header_fill
                )

                cell.font = (
                    accent_font
                    if sheet_name == "Key Findings"
                    else header_font
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

                cell.border = (
                    thin_border
                )

            data_start_row = (
                header_row + 1
            )

        else:
            data_start_row = 1

        for row_idx, row in enumerate(
            rows or [],
            start=data_start_row,
        ):

            if not isinstance(
                row,
                (
                    list,
                    tuple,
                ),
            ):
                row = [
                    row
                ]

            for col_idx, value in enumerate(
                row,
                start=1,
            ):

                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=_normalize_cell_value(
                        value
                    ),
                )

                cell.font = (
                    body_font
                )

                cell.border = (
                    thin_border
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        if sheet_name == "Executive Summary":

            for row_idx in range(
                data_start_row,
                data_start_row
                + len(
                    rows or []
                ),
            ):

                metric_cell = ws.cell(
                    row=row_idx,
                    column=1,
                )

                metric_cell.font = Font(
                    name="Calibri",
                    size=10,
                    bold=True,
                    color="0F172A",
                )

        max_column = (
            max(
                len(headers),
                max(
                    (
                        len(row)
                        for row in (
                            rows or []
                        )
                        if isinstance(
                            row,
                            (
                                list,
                                tuple,
                            ),
                        )
                    ),
                    default=0,
                ),
            )
            if headers
            else max(
                (
                    len(row)
                    for row in (
                        rows or []
                    )
                    if isinstance(
                        row,
                        (
                            list,
                            tuple,
                        ),
                    )
                ),
                default=1,
            )
        )

        for col_idx in range(
            1,
            max_column + 1,
        ):

            max_length = 0

            for row_idx in range(
                1,
                ws.max_row + 1,
            ):

                value = ws.cell(
                    row=row_idx,
                    column=col_idx,
                ).value

                if value is None:
                    continue

                value_length = len(
                    str(value)
                )

                max_length = max(
                    max_length,
                    min(
                        value_length,
                        70,
                    ),
                )

            ws.column_dimensions[
                get_column_letter(
                    col_idx
                )
            ].width = max(
                12,
                min(
                    max_length + 3,
                    70,
                ),
            )

        if ws.max_row > 1:

            ws.freeze_panes = (
                "A4"
                if sheet_name == "Executive Summary"
                else "A2"
            )

        ws.sheet_view.showGridLines = False

        if headers:

            filter_row = (
                3
                if sheet_name == "Executive Summary"
                else 1
            )

            if ws.max_row >= filter_row:

                ws.auto_filter.ref = (
                    f"A{filter_row}:"
                    f"{get_column_letter(max_column)}"
                    f"{ws.max_row}"
                )

    if "Executive Summary" in wb.sheetnames:

        wb._sheets.insert(
            0,
            wb._sheets.pop(
                wb.sheetnames.index(
                    "Executive Summary"
                )
            ),
        )

        wb.active = 0

    wb.save(
        str(path)
    )

    if not path.exists():
        raise RuntimeError(
            "XLSX workbook was not created."
        )

    file_size = path.stat().st_size

    if file_size <= 0:
        raise RuntimeError(
            "Generated XLSX workbook is empty."
        )

    verified, msg = verify_artifact(
        path,
        "xlsx",
    )

    if not verified:
        raise RuntimeError(
            f"XLSX creation failed verification: {msg}"
        )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": ".xlsx",
        "size_bytes": file_size,
        "created": True,
        "verification": msg,
        "workspace": "NOVA",
        "tool": "spreadsheet_writer",
        "report_type": (
            "spreadsheet_analysis_report"
            if source_file_path
            else "xlsx_workbook"
        ),
        "source_file": (
            source_file_path
            if source_file_path
            else None
        ),
    }


# ---------------------------------------------------------------------------
# CSV WRITER
# ---------------------------------------------------------------------------

def csv_writer_handler(
    file_path: str,
    headers: List[str] = None,
    rows: List[List[Any]] = None,
    content: str = "",
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():
        raise ValueError(
            "csv_writer requires a file_path."
        )

    path = _resolve_workspace_file(
        file_path
    )

    if path.suffix.lower() != ".csv":
        path = path.with_suffix(
            ".csv"
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):
        raise ValueError(
            "csv_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():
        raise ValueError(
            "csv_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        writer = csv.writer(
            f
        )

        if headers:
            writer.writerow(
                headers
            )

        if rows:

            for row in rows:
                writer.writerow(
                    row
                )

        elif content:

            f.write(
                content
            )

    verified, msg = verify_artifact(
        path,
        "csv",
    )

    if not verified:
        raise RuntimeError(
            f"CSV creation failed verification: {msg}"
        )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": ".csv",
        "size_bytes": path.stat().st_size,
        "created": True,
        "verification": msg,
        "workspace": "NOVA",
        "tool": "csv_writer",
    }


# ===========================================================================
# PPTX WRITER
# ===========================================================================

# ---------------------------------------------------------------------------
# NOVA PRESENTATION PALETTE
# ---------------------------------------------------------------------------

NOVA_BG = RGBColor(
    8,
    11,
    16,
)

NOVA_PANEL = RGBColor(
    15,
    23,
    32,
)

NOVA_PANEL_LIGHT = RGBColor(
    22,
    32,
    43,
)

NOVA_CYAN = RGBColor(
    0,
    240,
    255,
)

NOVA_BLUE = RGBColor(
    56,
    189,
    248,
)

NOVA_WHITE = RGBColor(
    241,
    245,
    249,
)

NOVA_MUTED = RGBColor(
    148,
    163,
    184,
)

NOVA_LINE = RGBColor(
    51,
    65,
    85,
)

NOVA_DARK = RGBColor(
    2,
    6,
    23,
)


# ---------------------------------------------------------------------------
# PPTX TEXT HELPERS
# ---------------------------------------------------------------------------

def _pptx_clean_text(
    text: Any,
) -> str:

    if text is None:
        return ""

    value = str(
        text
    ).strip()

    replacements = (
        (r"\*\*", "**"),
        (r"\_\_", "__"),
        (r"\*", "*"),
        (r"\_", "_"),
        (r"\`", "`"),
        (r"\#", "#"),
        (r"\-", "-"),
    )

    for old, new in replacements:
        value = value.replace(
            old,
            new,
        )

    value = re.sub(
        r"^\s*#{1,6}\s+",
        "",
        value,
    )

    return value.strip(
        "*_` "
    )


def _pptx_inline_segments(
    text: str,
) -> List[Dict[str, Any]]:

    normalized = _pptx_clean_text(
        text
    )

    if not normalized:
        return []

    pattern = re.compile(
        r"("
        r"\*\*(.+?)\*\*"
        r"|__(.+?)__"
        r"|(?<!\*)\*(?!\s)(.+?)(?<!\s)\*"
        r"|(?<!_)_(?!\s)(.+?)(?<!\s)_"
        r"|`(.+?)`"
        r")",
        flags=re.DOTALL,
    )

    segments: List[Dict[str, Any]] = []

    cursor = 0

    for match in pattern.finditer(
        normalized
    ):

        start, end = match.span()

        if start > cursor:

            segments.append(
                {
                    "text": normalized[
                        cursor:start
                    ],
                    "bold": False,
                    "italic": False,
                    "code": False,
                }
            )

        matched = match.group(0)

        if (
            matched.startswith("**")
            and matched.endswith("**")
        ):

            segments.append(
                {
                    "text": matched[
                        2:-2
                    ],
                    "bold": True,
                    "italic": False,
                    "code": False,
                }
            )

        elif (
            matched.startswith("__")
            and matched.endswith("__")
        ):

            segments.append(
                {
                    "text": matched[
                        2:-2
                    ],
                    "bold": True,
                    "italic": False,
                    "code": False,
                }
            )

        elif (
            matched.startswith("*")
            and matched.endswith("*")
        ):

            segments.append(
                {
                    "text": matched[
                        1:-1
                    ],
                    "bold": False,
                    "italic": True,
                    "code": False,
                }
            )

        elif (
            matched.startswith("_")
            and matched.endswith("_")
        ):

            segments.append(
                {
                    "text": matched[
                        1:-1
                    ],
                    "bold": False,
                    "italic": True,
                    "code": False,
                }
            )

        elif (
            matched.startswith("`")
            and matched.endswith("`")
        ):

            segments.append(
                {
                    "text": matched[
                        1:-1
                    ],
                    "bold": False,
                    "italic": False,
                    "code": True,
                }
            )

        cursor = end

    if cursor < len(normalized):

        segments.append(
            {
                "text": normalized[
                    cursor:
                ],
                "bold": False,
                "italic": False,
                "code": False,
            }
        )

    return segments


def _pptx_add_text_runs(
    paragraph: Any,
    text: str,
    font_size: float,
    font_color: RGBColor = NOVA_WHITE,
    font_name: str = "Aptos",
    bold_default: bool = False,
) -> None:

    segments = _pptx_inline_segments(
        text
    )

    if not segments:

        run = paragraph.add_run()
        run.text = ""

        return

    for segment in segments:

        run = paragraph.add_run()

        run.text = segment[
            "text"
        ]

        run.font.name = (
            "Consolas"
            if segment.get(
                "code"
            )
            else font_name
        )

        run.font.size = PPTXPt(
            9.5
            if segment.get(
                "code"
            )
            else font_size
        )

        run.font.bold = (
            bool(
                segment.get(
                    "bold",
                    False,
                )
            )
            or bold_default
        )

        run.font.italic = bool(
            segment.get(
                "italic",
                False,
            )
        )

        run.font.color.rgb = (
            NOVA_BLUE
            if segment.get(
                "code"
            )
            else font_color
        )


# ---------------------------------------------------------------------------
# PPTX BACKGROUND / FOOTER
# ---------------------------------------------------------------------------

def _pptx_add_background(
    slide: Any,
) -> None:

    fill = slide.background.fill

    fill.solid()

    fill.fore_color.rgb = NOVA_BG

    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        PPTXInches(13.333),
        PPTXInches(0.055),
    )

    accent_line.fill.solid()

    accent_line.fill.fore_color.rgb = (
        NOVA_CYAN
    )

    accent_line.line.fill.background()

    lower_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PPTXInches(0.55),
        PPTXInches(7.23),
        PPTXInches(12.23),
        PPTXInches(0.012),
    )

    lower_line.fill.solid()

    lower_line.fill.fore_color.rgb = (
        NOVA_LINE
    )

    lower_line.line.fill.background()


def _pptx_add_footer(
    slide: Any,
    slide_number: int,
) -> None:

    footer = slide.shapes.add_textbox(
        PPTXInches(0.6),
        PPTXInches(7.27),
        PPTXInches(10.8),
        PPTXInches(0.25),
    )

    tf = footer.text_frame

    tf.clear()

    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    paragraph = tf.paragraphs[0]

    paragraph.alignment = (
        PP_ALIGN.LEFT
    )

    run = paragraph.add_run()

    run.text = (
        "NOVA  /  SOVEREIGN INDUSTRIAL INTELLIGENCE"
    )

    run.font.name = "Aptos"
    run.font.size = PPTXPt(
        7.5
    )

    run.font.bold = True

    run.font.color.rgb = NOVA_MUTED

    page_box = slide.shapes.add_textbox(
        PPTXInches(11.5),
        PPTXInches(7.22),
        PPTXInches(1.2),
        PPTXInches(0.3),
    )

    page_tf = page_box.text_frame

    page_tf.clear()

    page_paragraph = page_tf.paragraphs[0]

    page_paragraph.alignment = (
        PP_ALIGN.RIGHT
    )

    page_run = page_paragraph.add_run()

    page_run.text = (
        f"{slide_number:02d}"
    )

    page_run.font.name = "Aptos"
    page_run.font.size = PPTXPt(
        8
    )

    page_run.font.bold = True

    page_run.font.color.rgb = NOVA_CYAN


# ---------------------------------------------------------------------------
# PPTX COMPONENTS
# ---------------------------------------------------------------------------

def _pptx_add_title_text(
    slide: Any,
    title: str,
    x: float = 0.8,
    y: float = 0.55,
    width: float = 11.7,
    height: float = 0.8,
    size: float = 27,
    color: RGBColor = NOVA_WHITE,
) -> None:

    box = slide.shapes.add_textbox(
        PPTXInches(x),
        PPTXInches(y),
        PPTXInches(width),
        PPTXInches(height),
    )

    tf = box.text_frame

    tf.clear()

    tf.word_wrap = True

    tf.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    paragraph = tf.paragraphs[0]

    paragraph.alignment = (
        PP_ALIGN.LEFT
    )

    _pptx_add_text_runs(
        paragraph=paragraph,
        text=title,
        font_size=size,
        font_color=color,
        font_name="Aptos Display",
        bold_default=True,
    )


def _pptx_add_panel(
    slide: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    fill_color: RGBColor = NOVA_PANEL,
    line_color: RGBColor = NOVA_LINE,
) -> Any:

    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PPTXInches(x),
        PPTXInches(y),
        PPTXInches(width),
        PPTXInches(height),
    )

    panel.fill.solid()

    panel.fill.fore_color.rgb = (
        fill_color
    )

    panel.line.color.rgb = (
        line_color
    )

    panel.line.width = PPTXPt(
        0.8
    )

    return panel


def _pptx_add_accent_chip(
    slide: Any,
    text: str,
    x: float,
    y: float,
    width: float = 2.0,
) -> None:

    chip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PPTXInches(x),
        PPTXInches(y),
        PPTXInches(width),
        PPTXInches(0.34),
    )

    chip.fill.solid()

    chip.fill.fore_color.rgb = (
        NOVA_CYAN
    )

    chip.line.fill.background()

    tf = chip.text_frame

    tf.clear()

    tf.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    paragraph = tf.paragraphs[0]

    paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    run = paragraph.add_run()

    run.text = _pptx_clean_text(
        text
    ).upper()

    run.font.name = "Aptos"
    run.font.size = PPTXPt(
        7.5
    )

    run.font.bold = True

    run.font.color.rgb = NOVA_DARK


def _pptx_add_bullet_block(
    slide: Any,
    bullets: List[str],
    x: float,
    y: float,
    width: float,
    height: float,
    font_size: float = 17,
    bullet_color: RGBColor = NOVA_CYAN,
    text_color: RGBColor = NOVA_WHITE,
) -> None:

    box = slide.shapes.add_textbox(
        PPTXInches(x),
        PPTXInches(y),
        PPTXInches(width),
        PPTXInches(height),
    )

    tf = box.text_frame

    tf.clear()

    tf.word_wrap = True

    tf.margin_left = PPTXInches(
        0.05
    )

    tf.margin_right = PPTXInches(
        0.05
    )

    tf.margin_top = PPTXInches(
        0.02
    )

    tf.margin_bottom = PPTXInches(
        0.02
    )

    for index, bullet in enumerate(
        bullets
    ):

        paragraph = (
            tf.paragraphs[0]
            if index == 0
            else tf.add_paragraph()
        )

        paragraph.text = ""

        paragraph.level = 0

        paragraph.space_after = PPTXPt(
            11
        )

        paragraph.line_spacing = 1.06

        bullet_run = paragraph.add_run()

        bullet_run.text = "◆"

        bullet_run.font.name = "Aptos"
        bullet_run.font.size = PPTXPt(
            max(
                10,
                font_size - 2,
            )
        )

        bullet_run.font.bold = True

        bullet_run.font.color.rgb = (
            bullet_color
        )

        cleaned_bullet = (
            _pptx_clean_text(
                bullet
            )
        )

        segments = _pptx_inline_segments(
            "  "
            + cleaned_bullet
        )

        if not segments:
            continue

        for segment in segments:

            run = paragraph.add_run()

            run.text = segment[
                "text"
            ]

            run.font.name = (
                "Consolas"
                if segment.get(
                    "code"
                )
                else "Aptos"
            )

            run.font.size = PPTXPt(
                9.5
                if segment.get(
                    "code"
                )
                else font_size
            )

            run.font.bold = bool(
                segment.get(
                    "bold",
                    False,
                )
            )

            run.font.italic = bool(
                segment.get(
                    "italic",
                    False,
                )
            )

            run.font.color.rgb = (
                NOVA_BLUE
                if segment.get("code")
                else text_color
            )


def _pptx_add_section_marker(
    slide: Any,
    text: str,
    x: float,
    y: float,
) -> None:

    marker = slide.shapes.add_textbox(
        PPTXInches(x),
        PPTXInches(y),
        PPTXInches(4.0),
        PPTXInches(0.32),
    )

    tf = marker.text_frame

    tf.clear()

    paragraph = tf.paragraphs[0]

    run = paragraph.add_run()

    run.text = (
        "/// "
        + _pptx_clean_text(
            text
        ).upper()
    )

    run.font.name = "Aptos"
    run.font.size = PPTXPt(
        8
    )

    run.font.bold = True

    run.font.color.rgb = NOVA_BLUE


# ---------------------------------------------------------------------------
# PPTX CONTENT PARSER
# ---------------------------------------------------------------------------

def _pptx_parse_content(
    content: str,
) -> List[Dict[str, Any]]:

    cleaned = _clean_document_text(
        content
    )

    if not cleaned:
        return []

    records: List[Dict[str, Any]] = []

    for raw_line in cleaned.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        if re.match(
            r"^\s*#{1,6}\s+",
            line,
        ):

            records.append(
                {
                    "type": "heading",
                    "text": _clean_heading_text(
                        line
                    ),
                }
            )

            continue

        number_match = re.match(
            r"^\s*\d+\.\s+(.+)$",
            line,
        )

        if number_match:

            records.append(
                {
                    "type": "numbered",
                    "text": number_match.group(
                        1
                    ).strip(),
                }
            )

            continue

        bullet_match = re.match(
            r"^\s*[-*+•·]\s+(.+)$",
            line,
        )

        if bullet_match:

            records.append(
                {
                    "type": "bullet",
                    "text": bullet_match.group(
                        1
                    ).strip(),
                }
            )

            continue

        if re.match(
            r"^\s*([-*_])(?:\s*\1){2,}\s*$",
            line,
        ):

            records.append(
                {
                    "type": "rule",
                    "text": "",
                }
            )

            continue

        records.append(
            {
                "type": "paragraph",
                "text": _pptx_clean_text(
                    line
                ),
            }
        )

    return records


def _pptx_group_records(
    records: List[Dict[str, Any]],
    max_bullets: int = 5,
    max_chars: int = 500,
) -> List[Dict[str, Any]]:

    groups: List[Dict[str, Any]] = []

    current_heading = "Overview"

    current_items: List[str] = []

    def flush():

        nonlocal current_items

        if not current_items:
            return

        groups.append(
            {
                "title": current_heading,
                "items": current_items[:],
            }
        )

        current_items = []

    for record in records:

        record_type = record.get(
            "type"
        )

        text = _pptx_clean_text(
            record.get(
                "text",
                "",
            )
        )

        if (
            not text
            and record_type != "rule"
        ):
            continue

        if record_type == "heading":

            flush()

            current_heading = (
                text
                or "Overview"
            )

            continue

        if record_type == "rule":

            flush()

            continue

        if record_type in {
            "bullet",
            "numbered",
        }:

            if (
                len(current_items)
                >= max_bullets
            ):
                flush()

            current_items.append(
                text
            )

            continue

        if (
            len(current_items)
            >= max_bullets
        ):
            flush()

        if len(text) > max_chars:

            text = (
                text[:max_chars]
                .rstrip()
                + "…"
            )

        current_items.append(
            text
        )

    flush()

    return groups


# ---------------------------------------------------------------------------
# FAST LOCAL VISUAL GENERATION
# ---------------------------------------------------------------------------

NOVA_IMAGE_GENERATION_ENABLED = (
    os.getenv(
        "NOVA_IMAGE_GENERATION_ENABLED",
        "true",
    )
    .strip()
    .lower()
    in {
        "1",
        "true",
        "yes",
        "on",
    }
)

NOVA_IMAGE_MODEL_ID = os.getenv(
    "NOVA_IMAGE_MODEL_ID",
    "stabilityai/sd-turbo",
).strip()

NOVA_IMAGE_LOCAL_FILES_ONLY = (
    os.getenv(
        "NOVA_IMAGE_LOCAL_FILES_ONLY",
        "false",
    )
    .strip()
    .lower()
    in {
        "1",
        "true",
        "yes",
        "on",
    }
)

# One-step SD-Turbo configuration.
NOVA_IMAGE_STEPS = int(
    os.getenv(
        "NOVA_IMAGE_STEPS",
        "1",
    )
)

# SD-Turbo is guidance-distilled.
NOVA_IMAGE_GUIDANCE = float(
    os.getenv(
        "NOVA_IMAGE_GUIDANCE",
        "0.0",
    )
)

# IMPORTANT:
# Diffusion happens at 256x256 for the 4 GB RTX 2050.
# The resulting image is upscaled to 512x512 afterwards.
NOVA_IMAGE_SIZE = int(
    os.getenv(
        "NOVA_IMAGE_SIZE",
        "256",
    )
)

# Final image size used by PowerPoint.
NOVA_IMAGE_OUTPUT_SIZE = int(
    os.getenv(
        "NOVA_IMAGE_OUTPUT_SIZE",
        "512",
    )
)

NOVA_IMAGE_DEVICE_MODE = (
    os.getenv(
        "NOVA_IMAGE_DEVICE_MODE",
        "auto",
    )
    .strip()
    .lower()
)

NOVA_IMAGE_DIRECT_GPU = (
    os.getenv(
        "NOVA_IMAGE_DIRECT_GPU",
        "true",
    )
    .strip()
    .lower()
    in {
        "1",
        "true",
        "yes",
        "on",
    }
)

_nova_image_pipeline = None
_nova_image_pipeline_lock = threading.Lock()
_nova_image_generation_lock = threading.Lock()


def _pptx_visual_profile(
    title: str,
    bullets: List[str],
) -> str:

    text = (
        _pptx_clean_text(
            title
        )
        + " "
        + " ".join(
            _pptx_clean_text(
                item
            )
            for item in (
                bullets or []
            )
        )
    ).lower()

    keyword_groups = {
        "computer_vision": (
            "computer vision",
            "vision",
            "visual inspection",
            "image recognition",
            "image analysis",
            "quality inspection",
            "defect detection",
            "camera",
            "visual",
            "ocr",
        ),
        "predictive_maintenance": (
            "predictive maintenance",
            "maintenance",
            "downtime",
            "machine health",
            "condition monitoring",
            "sensor",
            "failure prediction",
            "equipment",
            "machine failure",
        ),
        "machine_learning": (
            "machine learning",
            "deep learning",
            "neural network",
            "model",
            "prediction",
            "classification",
            "training",
            "algorithm",
        ),
        "cybersecurity": (
            "cybersecurity",
            "security",
            "secure",
            "protection",
            "threat",
            "risk",
            "data protection",
            "attack",
            "access control",
            "zero trust",
        ),
        "supply_chain": (
            "supply chain",
            "logistics",
            "inventory",
            "warehouse",
            "shipping",
            "distribution",
            "procurement",
        ),
        "data_analytics": (
            "data analytics",
            "analytics",
            "dashboard",
            "metrics",
            "kpi",
            "performance",
            "statistics",
            "decision support",
            "monitoring",
        ),
        "implementation": (
            "implementation",
            "deployment",
            "roadmap",
            "strategy",
            "integration",
            "phased",
            "adoption",
            "rollout",
        ),
        "future": (
            "future",
            "next generation",
            "innovation",
            "smart factory",
            "industry 4.0",
            "connected",
            "autonomous",
            "digital twin",
            "future factory",
        ),
        "benefits": (
            "benefit",
            "productivity",
            "efficiency",
            "quality improvement",
            "revenue",
            "cost reduction",
            "optimization",
            "growth",
        ),
        "challenges": (
            "challenge",
            "limitation",
            "roi",
            "difficulty",
            "barrier",
            "constraint",
        ),
        "industrial_ai": (
            "industrial",
            "factory",
            "manufacturing",
            "automation",
            "robot",
            "robotics",
            "production",
            "plant",
            "industry",
        ),
    }

    priority_order = [
        "computer_vision",
        "predictive_maintenance",
        "machine_learning",
        "cybersecurity",
        "supply_chain",
        "data_analytics",
        "implementation",
        "future",
        "challenges",
        "benefits",
        "industrial_ai",
    ]

    for profile in priority_order:

        if any(
            keyword in text
            for keyword in keyword_groups[
                profile
            ]
        ):
            return profile

    return "industrial_ai"


def _pptx_visual_prompt(
    title: str,
    bullets: List[str],
    profile: str,
    slide_index: int,
) -> str:
    """
    Build a compact CLIP-safe prompt.
    """

    clean_title = re.sub(
        r"\s+",
        " ",
        _pptx_clean_text(
            title
        ),
    ).strip()

    profile_descriptions = {
        "computer_vision": (
            "industrial inspection camera, "
            "factory quality control, machine vision"
        ),
        "predictive_maintenance": (
            "industrial machine sensors, "
            "predictive maintenance, factory equipment"
        ),
        "machine_learning": (
            "industrial AI, neural network visualization, "
            "intelligent manufacturing"
        ),
        "cybersecurity": (
            "industrial cybersecurity, "
            "secure factory network, protected control systems"
        ),
        "supply_chain": (
            "smart warehouse, industrial logistics, "
            "automated supply chain"
        ),
        "data_analytics": (
            "industrial analytics dashboard, "
            "smart factory monitoring"
        ),
        "implementation": (
            "engineers deploying AI in a factory, "
            "industrial technology integration"
        ),
        "future": (
            "near-future smart factory, "
            "connected machines, advanced robotics"
        ),
        "benefits": (
            "efficient automated factory, "
            "productive industrial robotics"
        ),
        "challenges": (
            "complex industrial operations, "
            "secure manufacturing environment"
        ),
        "industrial_ai": (
            "modern smart factory, "
            "industrial robots, automated production line"
        ),
    }

    visual_subject = profile_descriptions.get(
        profile,
        profile_descriptions[
            "industrial_ai"
        ],
    )

    bullet_context = [
        re.sub(
            r"\s+",
            " ",
            _pptx_clean_text(
                item
            ),
        ).strip()
        for item in (
            bullets or []
        )[:2]
        if _pptx_clean_text(
            item
        )
    ]

    context = (
        ", ".join(
            bullet_context
        )
        if bullet_context
        else clean_title
    )

    return (
        f"{visual_subject}. "
        f"{clean_title}. "
        f"{context}. "
        "photorealistic industrial photography, "
        "real machinery, cinematic factory lighting, "
        "wide composition, no text, no logo."
    )


def _load_nova_image_pipeline():
    """
    Lazily load and cache the diffusion pipeline.

    The model is loaded once per backend process.
    """

    global _nova_image_pipeline

    if not NOVA_IMAGE_GENERATION_ENABLED:
        return None

    if _nova_image_pipeline is not None:
        return _nova_image_pipeline

    with _nova_image_pipeline_lock:

        if _nova_image_pipeline is not None:
            return _nova_image_pipeline

        try:
            import torch

            from diffusers import (
                AutoPipelineForText2Image,
            )

        except ImportError as exc:

            raise RuntimeError(
                "Local image generation requires "
                "diffusers, accelerate, safetensors, and torch."
            ) from exc

        has_cuda = bool(
            torch.cuda.is_available()
        )

        dtype = (
            torch.float16
            if has_cuda
            else torch.float32
        )

        try:
            torch.backends.cuda.matmul.allow_tf32 = True
        except Exception:
            pass

        print(
            "[NOVA IMAGE] "
            f"loading {NOVA_IMAGE_MODEL_ID} | "
            f"cuda={has_cuda} | "
            f"dtype={dtype}"
        )

        try:

            pipe = AutoPipelineForText2Image.from_pretrained(
                NOVA_IMAGE_MODEL_ID,
                torch_dtype=dtype,
                variant="fp16" if has_cuda else None,
                local_files_only=(
                    NOVA_IMAGE_LOCAL_FILES_ONLY
                ),
            )

        except Exception:

            pipe = AutoPipelineForText2Image.from_pretrained(
                NOVA_IMAGE_MODEL_ID,
                torch_dtype=dtype,
                local_files_only=(
                    NOVA_IMAGE_LOCAL_FILES_ONLY
                ),
            )

        if has_cuda and NOVA_IMAGE_DIRECT_GPU:

            direct_gpu_loaded = False

            try:

                pipe = pipe.to(
                    "cuda"
                )

                direct_gpu_loaded = True

                print(
                    "[NOVA IMAGE] "
                    "direct CUDA mode enabled"
                )

            except Exception as direct_error:

                print(
                    "[NOVA IMAGE GPU FALLBACK] "
                    f"{type(direct_error).__name__}: "
                    f"{direct_error}"
                )

            if not direct_gpu_loaded:

                try:

                    pipe.enable_model_cpu_offload()

                    print(
                        "[NOVA IMAGE] "
                        "model CPU offload enabled"
                    )

                except Exception:

                    try:

                        pipe.enable_sequential_cpu_offload()

                        print(
                            "[NOVA IMAGE] "
                            "sequential CPU offload enabled"
                        )

                    except Exception:

                        pipe = pipe.to(
                            "cpu"
                        )

                        print(
                            "[NOVA IMAGE] "
                            "CPU mode enabled"
                        )

        elif has_cuda:

            try:

                pipe.enable_model_cpu_offload()

                print(
                    "[NOVA IMAGE] "
                    "model CPU offload enabled"
                )

            except Exception:

                try:

                    pipe.enable_sequential_cpu_offload()

                    print(
                        "[NOVA IMAGE] "
                        "sequential CPU offload enabled"
                    )

                except Exception:

                    pipe = pipe.to(
                        "cpu"
                    )

        else:

            pipe = pipe.to(
                "cpu"
            )

            print(
                "[NOVA IMAGE] "
                "CPU mode enabled"
            )

        try:

            pipe.set_progress_bar_config(
                disable=True
            )

        except Exception:
            pass

        try:
            if hasattr(
                pipe,
                "disable_attention_slicing",
            ):
                pipe.disable_attention_slicing()
        except Exception:
            pass

        # SD-Turbo does not need VAE slicing in this one-step path.
        try:
            if hasattr(
                pipe,
                "disable_vae_slicing",
            ):
                pipe.disable_vae_slicing()
        except Exception:
            pass

        _nova_image_pipeline = pipe

    return _nova_image_pipeline


def _pptx_generate_real_image_asset(
    title: str,
    bullets: List[str],
    slide_index: int,
) -> str:
    """
    Generate one topic-aware real image locally.

    Fast 4 GB GPU path:
    - SD-Turbo
    - one inference step
    - 256x256 diffusion
    - cached pipeline
    - cached individual images
    - 512x512 final output
    - no negative prompt
    - torch.inference_mode()
    """

    visual_profile = _pptx_visual_profile(
        title,
        bullets,
    )

    prompt = _pptx_visual_prompt(
        title=title,
        bullets=bullets,
        profile=visual_profile,
        slide_index=slide_index,
    )

    visual_root = (
        WORKSPACE_ROOT
        / "temp"
        / "ppt_generated_images"
    ).resolve()

    visual_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    clean_bullets = [
        _pptx_clean_text(
            item
        )
        for item in (
            bullets or []
        )
    ]

    # Version tag prevents old 512x512 cached assets from being reused.
    signature = (
        "NOVA_FAST_SD_TURBO_V2|"
        f"{title}|"
        f"{'|'.join(clean_bullets)}|"
        f"{slide_index}|"
        f"{NOVA_IMAGE_MODEL_ID}|"
        f"{NOVA_IMAGE_STEPS}|"
        f"{NOVA_IMAGE_SIZE}|"
        f"{NOVA_IMAGE_OUTPUT_SIZE}"
    )

    digest = hashlib.sha1(
        signature.encode(
            "utf-8"
        )
    ).hexdigest()[:16]

    output_path = (
        visual_root
        / f"generated_{digest}.png"
    )

    if (
        output_path.exists()
        and output_path.stat().st_size > 0
    ):
        return str(
            output_path
        )

    pipe = _load_nova_image_pipeline()

    if pipe is None:

        raise RuntimeError(
            "NOVA local image generation is disabled."
        )

    steps = max(
        1,
        min(
            NOVA_IMAGE_STEPS,
            4,
        ),
    )

    generation_size = max(
        128,
        min(
            NOVA_IMAGE_SIZE,
            512,
        ),
    )

    final_size = max(
        256,
        min(
            NOVA_IMAGE_OUTPUT_SIZE,
            1024,
        ),
    )

    generation_kwargs: Dict[str, Any] = {
        "prompt": prompt,
        "num_inference_steps": steps,
        "guidance_scale": max(
            0.0,
            min(
                NOVA_IMAGE_GUIDANCE,
                2.0,
            ),
        ),
        "width": generation_size,
        "height": generation_size,
    }

    try:

        import torch

        seed = (
            int(
                digest,
                16,
            )
            % (
                2**31 - 1
            )
        )

        generator_device = (
            "cuda"
            if torch.cuda.is_available()
            and NOVA_IMAGE_DIRECT_GPU
            else "cpu"
        )

        generator = torch.Generator(
            device=generator_device
        ).manual_seed(
            seed
        )

        generation_kwargs[
            "generator"
        ] = generator

    except Exception:
        pass

    started = datetime.now()

    with _nova_image_generation_lock:

        if (
            output_path.exists()
            and output_path.stat().st_size > 0
        ):
            return str(
                output_path
            )

        try:

            import torch

            with torch.inference_mode():

                result = pipe(
                    **generation_kwargs
                )

        except TypeError:

            generation_kwargs.pop(
                "guidance_scale",
                None,
            )

            try:

                import torch

                with torch.inference_mode():

                    result = pipe(
                        **generation_kwargs
                    )

            except Exception:

                result = pipe(
                    **generation_kwargs
                )

        except Exception:

            # Final compatibility fallback for unusual CPU-only
            # pipelines that do not cooperate with inference_mode.
            result = pipe(
                **generation_kwargs
            )

        if not result.images:

            raise RuntimeError(
                "Local image generation returned no image."
            )

        image = result.images[0]

        if image.mode != "RGB":

            image = image.convert(
                "RGB"
            )

        # Upscale only after diffusion.
        if image.size != (
            final_size,
            final_size,
        ):

            image = image.resize(
                (
                    final_size,
                    final_size,
                ),
                Image.Resampling.LANCZOS,
            )

        image.save(
            output_path,
            "PNG",
            optimize=True,
        )

        if (
            not output_path.exists()
            or output_path.stat().st_size <= 0
        ):

            raise RuntimeError(
                "Local image generation did not create a valid PNG."
            )

    elapsed = (
        datetime.now()
        - started
    ).total_seconds()

    print(
        "[NOVA IMAGE GENERATED] "
        f"{elapsed:.2f}s | "
        f"profile={visual_profile} | "
        f"diffusion={generation_size}x{generation_size} | "
        f"output={final_size}x{final_size}"
    )

    return str(
        output_path
    )


# ---------------------------------------------------------------------------
# FAST LOCAL FALLBACK VISUALS
# ---------------------------------------------------------------------------

def _pptx_draw_grid(
    draw: Any,
    width: int,
    height: int,
) -> None:

    grid_color = (
        22,
        39,
        52,
    )

    for x in range(
        0,
        width,
        70,
    ):

        draw.line(
            [
                (x, 0),
                (x, height),
            ],
            fill=grid_color,
            width=1,
        )

    for y in range(
        0,
        height,
        70,
    ):

        draw.line(
            [
                (0, y),
                (width, y),
            ],
            fill=grid_color,
            width=1,
        )


def _pptx_draw_node(
    draw: Any,
    x: int,
    y: int,
    radius: int = 10,
    color=(0, 240, 255),
) -> None:

    for spread in (
        radius * 3,
        radius * 2,
        radius,
    ):

        alpha_scale = (
            0.16
            if spread == radius * 3
            else 0.26
            if spread == radius * 2
            else 1.0
        )

        c = tuple(
            int(
                channel * alpha_scale
            )
            for channel in color
        )

        draw.ellipse(
            [
                x - spread,
                y - spread,
                x + spread,
                y + spread,
            ],
            fill=c,
        )

    draw.ellipse(
        [
            x - radius,
            y - radius,
            x + radius,
            y + radius,
        ],
        fill=color,
    )


def _pptx_draw_circuit(
    draw: Any,
    points: List[tuple],
    color=(56, 189, 248),
    width: int = 5,
) -> None:

    if len(points) < 2:
        return

    draw.line(
        points,
        fill=color,
        width=width,
    )

    for x, y in points:

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=7,
            color=color,
        )


def _pptx_draw_factory_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    white = (
        210,
        225,
        235,
    )

    panel = (
        19,
        32,
        43,
    )

    draw.rounded_rectangle(
        [
            170,
            360,
            940,
            690,
        ],
        radius=28,
        fill=panel,
        outline=(45, 67, 82),
        width=4,
    )

    draw.polygon(
        [
            (170, 360),
            (300, 275),
            (460, 360),
        ],
        fill=(24, 40, 53),
        outline=blue,
    )

    draw.polygon(
        [
            (455, 360),
            (600, 250),
            (790, 360),
        ],
        fill=(22, 37, 51),
        outline=blue,
    )

    draw.polygon(
        [
            (760, 360),
            (845, 295),
            (940, 360),
        ],
        fill=(21, 35, 48),
        outline=blue,
    )

    for row in range(2):

        for col in range(5):

            x = 220 + col * 135
            y = 420 + row * 90

            draw.rounded_rectangle(
                [
                    x,
                    y,
                    x + 70,
                    y + 46,
                ],
                radius=7,
                fill=(8, 19, 27),
                outline=(0, 120, 145),
                width=2,
            )

            draw.line(
                [
                    (x + 35, y),
                    (x + 35, y + 46),
                ],
                fill=(42, 79, 97),
                width=1,
            )

    draw.rounded_rectangle(
        [
            90,
            650,
            1020,
            755,
        ],
        radius=22,
        fill=(11, 20, 27),
        outline=cyan,
        width=4,
    )

    for x in range(
        130,
        980,
        70,
    ):

        draw.ellipse(
            [
                x,
                680,
                x + 38,
                718,
            ],
            outline=(63, 89, 105),
            width=2,
        )

    draw.line(
        [
            (1080, 640),
            (1080, 510),
            (1010, 445),
            (960, 505),
        ],
        fill=cyan,
        width=16,
    )

    draw.ellipse(
        [
            1045,
            605,
            1115,
            675,
        ],
        fill=(18, 34, 45),
        outline=cyan,
        width=4,
    )

    draw.ellipse(
        [
            985,
            420,
            1045,
            480,
        ],
        fill=(18, 34, 45),
        outline=blue,
        width=4,
    )

    draw.rectangle(
        [
            950,
            480,
            1002,
            532,
        ],
        fill=(25, 43, 56),
        outline=white,
        width=2,
    )

    points = [
        (1030, 120),
        (1140, 90),
        (1240, 160),
        (1160, 230),
        (1030, 210),
    ]

    _pptx_draw_circuit(
        draw,
        points,
        color=cyan,
        width=4,
    )


def _pptx_draw_ml_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    purple = (
        129,
        140,
        248,
    )

    layers = [
        [150, 290, 430, 570, 730],
        [260, 380, 500, 620],
        [180, 340, 500, 660, 820],
        [280, 470, 660],
    ]

    coordinates = []

    for xs in layers:

        layer_nodes = []

        for idx, x in enumerate(
            xs
        ):

            y = (
                170
                + (
                    idx
                    * (
                        520
                        / max(
                            1,
                            len(xs) - 1,
                        )
                    )
                )
            )

            if len(xs) == 1:
                y = 400

            layer_nodes.append(
                (
                    int(x),
                    int(y),
                )
            )

        coordinates.append(
            layer_nodes
        )

    for layer_index in range(
        len(coordinates) - 1
    ):

        for start in coordinates[
            layer_index
        ]:

            for end in coordinates[
                layer_index + 1
            ]:

                draw.line(
                    [
                        start,
                        end,
                    ],
                    fill=(36, 64, 82),
                    width=3,
                )

    for layer_index, nodes in enumerate(
        coordinates
    ):

        node_color = (
            cyan
            if layer_index % 2 == 0
            else purple
        )

        for x, y in nodes:

            _pptx_draw_node(
                draw,
                x,
                y,
                radius=13,
                color=node_color,
            )

    bars = [
        220,
        340,
        280,
        410,
        370,
        455,
    ]

    for idx, value in enumerate(
        bars
    ):

        x = 1010 + idx * 38

        draw.rounded_rectangle(
            [
                x,
                570 - value / 3,
                x + 22,
                570,
            ],
            radius=8,
            fill=blue
            if idx % 2
            else cyan,
        )


def _pptx_draw_vision_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    white = (
        210,
        225,
        235,
    )

    draw.rounded_rectangle(
        [
            140,
            250,
            730,
            560,
        ],
        radius=38,
        fill=(18, 31, 42),
        outline=(54, 83, 102),
        width=5,
    )

    draw.rectangle(
        [
            190,
            190,
            420,
            270,
        ],
        fill=(19, 33, 44),
        outline=blue,
        width=4,
    )

    draw.ellipse(
        [
            310,
            315,
            560,
            565,
        ],
        fill=(8, 15, 21),
        outline=cyan,
        width=8,
    )

    draw.ellipse(
        [
            355,
            360,
            515,
            520,
        ],
        fill=(20, 43, 58),
        outline=blue,
        width=5,
    )

    draw.ellipse(
        [
            395,
            400,
            475,
            480,
        ],
        fill=(0, 120, 155),
    )

    draw.rounded_rectangle(
        [
            780,
            170,
            1240,
            640,
        ],
        radius=24,
        fill=(11, 21, 29),
        outline=(34, 65, 82),
        width=4,
    )

    draw.rounded_rectangle(
        [
            875,
            280,
            1130,
            520,
        ],
        radius=22,
        fill=(22, 38, 49),
        outline=white,
        width=4,
    )

    boxes = [
        (850, 235, 950, 330),
        (1060, 390, 1165, 500),
        (920, 445, 1030, 555),
    ]

    for box in boxes:

        draw.rectangle(
            box,
            outline=cyan,
            width=5,
        )

    draw.line(
        [
            (730, 380),
            (840, 320),
        ],
        fill=(64, 141, 168),
        width=4,
    )

    draw.line(
        [
            (730, 405),
            (835, 430),
        ],
        fill=(64, 141, 168),
        width=4,
    )

    draw.line(
        [
            (730, 430),
            (845, 525),
        ],
        fill=(64, 141, 168),
        width=4,
    )


def _pptx_draw_maintenance_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    draw.rounded_rectangle(
        [
            120,
            410,
            790,
            670,
        ],
        radius=28,
        fill=(18, 31, 42),
        outline=(54, 83, 102),
        width=5,
    )

    for x in (
        220,
        390,
        560,
    ):

        draw.ellipse(
            [
                x,
                285,
                x + 150,
                435,
            ],
            fill=(26, 46, 60),
            outline=blue,
            width=5,
        )

        draw.ellipse(
            [
                x + 32,
                317,
                x + 118,
                403,
            ],
            fill=(7, 15, 21),
            outline=cyan,
            width=4,
        )

        draw.line(
            [
                (x + 75, 435),
                (x + 75, 585),
            ],
            fill=blue,
            width=12,
        )

    for x, y in (
        (235, 240),
        (470, 220),
        (675, 260),
    ):

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=14,
            color=cyan,
        )

    wave = []

    for x in range(
        850,
        1245,
        12,
    ):

        y = (
            460
            + int(
                55
                * (
                    (
                        (x - 850)
                        % 95
                    )
                    / 95
                )
            )
        )

        wave.append(
            (
                x,
                y,
            )
        )

    draw.line(
        wave,
        fill=cyan,
        width=5,
    )


def _pptx_draw_cybersecurity_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    shield = [
        (690, 150),
        (940, 205),
        (900, 510),
        (690, 690),
        (480, 510),
        (440, 205),
    ]

    draw.polygon(
        shield,
        fill=(14, 32, 43),
        outline=cyan,
    )

    draw.line(
        shield + [shield[0]],
        fill=blue,
        width=7,
    )

    draw.rounded_rectangle(
        [
            580,
            355,
            800,
            550,
        ],
        radius=25,
        fill=(10, 22, 30),
        outline=cyan,
        width=5,
    )

    draw.arc(
        [
            620,
            285,
            760,
            440,
        ],
        start=180,
        end=360,
        fill=blue,
        width=12,
    )

    draw.rectangle(
        [
            677,
            425,
            705,
            482,
        ],
        fill=cyan,
    )

    for x, y in (
        (150, 210),
        (260, 420),
        (175, 620),
        (1090, 235),
        (1190, 430),
        (1085, 625),
    ):

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=12,
            color=blue,
        )

        draw.line(
            [
                (x, y),
                (690, 420),
            ],
            fill=(42, 76, 97),
            width=3,
        )


def _pptx_draw_supply_chain_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    route = [
        (145, 580),
        (350, 380),
        (590, 520),
        (830, 285),
        (1100, 455),
        (1240, 240),
    ]

    draw.line(
        route,
        fill=(45, 91, 112),
        width=18,
    )

    draw.line(
        route,
        fill=cyan,
        width=4,
    )

    for x, y in route:

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=13,
            color=blue,
        )

    draw.polygon(
        [
            (180, 270),
            (390, 155),
            (600, 270),
        ],
        fill=(20, 36, 48),
        outline=blue,
    )

    draw.rectangle(
        [
            180,
            270,
            600,
            550,
        ],
        fill=(15, 27, 37),
        outline=(48, 74, 89),
        width=4,
    )

    for row in range(2):

        for col in range(3):

            x = 230 + col * 105
            y = 320 + row * 100

            draw.rectangle(
                [
                    x,
                    y,
                    x + 75,
                    y + 65,
                ],
                fill=(30, 54, 67),
                outline=cyan,
                width=2,
            )

    draw.rounded_rectangle(
        [
            860,
            500,
            1180,
            650,
        ],
        radius=22,
        fill=(20, 36, 47),
        outline=blue,
        width=4,
    )


def _pptx_draw_data_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    draw.rounded_rectangle(
        [
            120,
            110,
            1240,
            690,
        ],
        radius=28,
        fill=(12, 23, 31),
        outline=(46, 74, 91),
        width=4,
    )

    for index in range(4):

        x = 165 + index * 255

        draw.rounded_rectangle(
            [
                x,
                160,
                x + 220,
                260,
            ],
            radius=18,
            fill=(19, 34, 45),
            outline=(38, 70, 88),
            width=2,
        )

        draw.rectangle(
            [
                x + 20,
                230,
                x + 110,
                242,
            ],
            fill=(
                cyan
                if index % 2 == 0
                else blue
            ),
        )

    for index, value in enumerate(
        [
            150,
            240,
            205,
            310,
            275,
            355,
            320,
        ]
    ):

        x = 190 + index * 100

        draw.rounded_rectangle(
            [
                x,
                590 - value / 1.5,
                x + 48,
                590,
            ],
            radius=10,
            fill=(
                cyan
                if index % 2 == 0
                else blue
            ),
        )

    points = [
        (180, 520),
        (300, 465),
        (420, 495),
        (545, 390),
        (665, 420),
        (790, 315),
        (915, 345),
        (1040, 245),
        (1160, 275),
    ]

    draw.line(
        points,
        fill=cyan,
        width=6,
    )

    for x, y in points:

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=7,
            color=cyan,
        )


def _pptx_draw_implementation_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    steps = [
        (170, 450),
        (420, 310),
        (680, 470),
        (940, 300),
        (1190, 445),
    ]

    draw.line(
        steps,
        fill=(45, 86, 106),
        width=14,
    )

    draw.line(
        steps,
        fill=cyan,
        width=5,
    )

    for index, (x, y) in enumerate(
        steps
    ):

        draw.rounded_rectangle(
            [
                x - 70,
                y - 70,
                x + 70,
                y + 70,
            ],
            radius=24,
            fill=(16, 31, 41),
            outline=blue,
            width=4,
        )

        draw.ellipse(
            [
                x - 28,
                y - 28,
                x + 28,
                y + 28,
            ],
            fill=(
                cyan
                if index % 2 == 0
                else blue
            ),
        )


def _pptx_draw_future_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    purple = (
        129,
        140,
        248,
    )

    draw.rounded_rectangle(
        [
            470,
            285,
            910,
            575,
        ],
        radius=34,
        fill=(17, 33, 43),
        outline=cyan,
        width=5,
    )

    nodes = [
        (180, 210, cyan),
        (370, 115, blue),
        (650, 80, purple),
        (980, 115, cyan),
        (1170, 230, blue),
        (1110, 550, purple),
        (880, 700, cyan),
        (560, 715, blue),
        (220, 560, purple),
    ]

    for x, y, color in nodes:

        draw.line(
            [
                (690, 430),
                (x, y),
            ],
            fill=(50, 90, 109),
            width=4,
        )

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=15,
            color=color,
        )

    for radius in (
        90,
        125,
        165,
    ):

        draw.ellipse(
            [
                690 - radius,
                430 - radius,
                690 + radius,
                430 + radius,
            ],
            outline=(
                cyan
                if radius == 90
                else blue
                if radius == 125
                else purple
            ),
            width=3,
        )


def _pptx_draw_benefits_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    purple = (
        129,
        140,
        248,
    )

    bars = [
        180,
        260,
        345,
        440,
        560,
    ]

    base_y = 690

    for idx, value in enumerate(
        bars
    ):

        x = 170 + idx * 190

        draw.rounded_rectangle(
            [
                x,
                base_y - value,
                x + 110,
                base_y,
            ],
            radius=18,
            fill=(
                cyan
                if idx < 2
                else blue
                if idx < 4
                else purple
            ),
        )

    points = [
        (170, 595),
        (355, 520),
        (545, 455),
        (735, 360),
        (925, 255),
        (1110, 165),
    ]

    draw.line(
        points,
        fill=cyan,
        width=7,
    )

    for x, y in points:

        _pptx_draw_node(
            draw,
            x,
            y,
            radius=9,
            color=cyan,
        )


def _pptx_draw_challenges_scene(
    draw: Any,
    width: int,
    height: int,
) -> None:

    cyan = (
        0,
        240,
        255,
    )

    blue = (
        56,
        189,
        248,
    )

    warning = (
        244,
        114,
        182,
    )

    pillar_xs = [
        190,
        550,
        910,
    ]

    pillar_heights = [
        300,
        420,
        350,
    ]

    for idx, x in enumerate(
        pillar_xs
    ):

        height_value = pillar_heights[
            idx
        ]

        draw.rounded_rectangle(
            [
                x,
                680 - height_value,
                x + 220,
                680,
            ],
            radius=22,
            fill=(18, 31, 41),
            outline=(
                cyan
                if idx == 0
                else blue
                if idx == 1
                else warning
            ),
            width=5,
        )

    _pptx_draw_circuit(
        draw,
        [
            (1100, 220),
            (1160, 310),
            (1075, 400),
            (1185, 495),
        ],
        color=warning,
        width=5,
    )


def _pptx_generate_fallback_visual_asset(
    title: str,
    bullets: List[str],
    slide_index: int,
) -> str:
    """
    Deterministic local fallback.

    This path is extremely fast and remains useful when:
    - image generation is disabled
    - CUDA is unavailable
    - diffusion model loading fails
    """

    clean_title = _pptx_clean_text(
        title
    )

    clean_bullets = [
        _pptx_clean_text(
            item
        )
        for item in (
            bullets or []
        )
        if _pptx_clean_text(
            item
        )
    ]

    signature = (
        f"{clean_title}|"
        f"{'|'.join(clean_bullets)}|"
        f"{slide_index}"
    )

    digest = hashlib.sha1(
        signature.encode(
            "utf-8"
        )
    ).hexdigest()[:16]

    visual_root = (
        WORKSPACE_ROOT
        / "temp"
        / "ppt_visuals"
    ).resolve()

    visual_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_path = (
        visual_root
        / f"fallback_{digest}.png"
    )

    if (
        output_path.exists()
        and output_path.stat().st_size > 0
    ):
        return str(
            output_path
        )

    width = 1400
    height = 850

    image = Image.new(
        "RGB",
        (
            width,
            height,
        ),
        (
            8,
            11,
            16,
        ),
    )

    draw = ImageDraw.Draw(
        image
    )

    _pptx_draw_grid(
        draw,
        width,
        height,
    )

    profile = _pptx_visual_profile(
        clean_title,
        clean_bullets,
    )

    if profile == "computer_vision":

        _pptx_draw_vision_scene(
            draw,
            width,
            height,
        )

    elif profile == "predictive_maintenance":

        _pptx_draw_maintenance_scene(
            draw,
            width,
            height,
        )

    elif profile == "machine_learning":

        _pptx_draw_ml_scene(
            draw,
            width,
            height,
        )

    elif profile == "cybersecurity":

        _pptx_draw_cybersecurity_scene(
            draw,
            width,
            height,
        )

    elif profile == "supply_chain":

        _pptx_draw_supply_chain_scene(
            draw,
            width,
            height,
        )

    elif profile == "data_analytics":

        _pptx_draw_data_scene(
            draw,
            width,
            height,
        )

    elif profile == "implementation":

        _pptx_draw_implementation_scene(
            draw,
            width,
            height,
        )

    elif profile == "future":

        _pptx_draw_future_scene(
            draw,
            width,
            height,
        )

    elif profile == "benefits":

        _pptx_draw_benefits_scene(
            draw,
            width,
            height,
        )

    elif profile == "challenges":

        _pptx_draw_challenges_scene(
            draw,
            width,
            height,
        )

    else:

        _pptx_draw_factory_scene(
            draw,
            width,
            height,
        )

    image.save(
        output_path,
        "PNG",
        optimize=True,
    )

    return str(
        output_path
    )


def _pptx_generate_visual_asset(
    title: str,
    bullets: List[str],
    slide_index: int,
) -> str:
    """
    Generate one topic-aware visual.

    Priority:
    1. cached real local image
    2. fast SD-Turbo image
    3. deterministic vector fallback
    """

    clean_bullets = [
        _pptx_clean_text(
            item
        )
        for item in (
            bullets or []
        )
    ]

    signature = (
        "NOVA_FAST_SD_TURBO_V2|"
        f"{title}|"
        f"{'|'.join(clean_bullets)}|"
        f"{slide_index}|"
        f"{NOVA_IMAGE_MODEL_ID}|"
        f"{NOVA_IMAGE_STEPS}|"
        f"{NOVA_IMAGE_SIZE}|"
        f"{NOVA_IMAGE_OUTPUT_SIZE}"
    )

    digest = hashlib.sha1(
        signature.encode(
            "utf-8"
        )
    ).hexdigest()[:16]

    generated_path = (
        WORKSPACE_ROOT
        / "temp"
        / "ppt_generated_images"
        / f"generated_{digest}.png"
    ).resolve()

    if (
        generated_path.exists()
        and generated_path.stat().st_size > 0
    ):
        return str(
            generated_path
        )

    if NOVA_IMAGE_GENERATION_ENABLED:

        started = datetime.now()

        try:

            image_path = _pptx_generate_real_image_asset(
                title=title,
                bullets=bullets,
                slide_index=slide_index,
            )

            elapsed = (
                datetime.now()
                - started
            ).total_seconds()

            print(
                "[NOVA FAST IMAGE] "
                f"{elapsed:.2f}s | "
                f"{_pptx_visual_profile(title, bullets)} | "
                f"{image_path}"
            )

            return image_path

        except Exception as exc:

            elapsed = (
                datetime.now()
                - started
            ).total_seconds()

            print(
                "[NOVA IMAGE FALLBACK] "
                f"{elapsed:.2f}s | "
                f"{type(exc).__name__}: {exc}"
            )

    return _pptx_generate_fallback_visual_asset(
        title=title,
        bullets=bullets,
        slide_index=slide_index,
    )


# ---------------------------------------------------------------------------
# PPTX IMAGE PLACEMENT
# ---------------------------------------------------------------------------

def _pptx_add_picture_cover(
    slide: Any,
    image_path: str,
    x: float,
    y: float,
    width: float,
    height: float,
) -> Any:

    resolved = _resolve_workspace_file(
        image_path
    )

    if not resolved.exists():
        raise FileNotFoundError(
            f"Presentation image not found: {image_path}"
        )

    target_width = float(
        width
    )

    target_height = float(
        height
    )

    with Image.open(
        resolved
    ) as image:

        image_width, image_height = (
            image.size
        )

    if (
        image_width <= 0
        or image_height <= 0
    ):
        raise ValueError(
            "Presentation image has invalid dimensions."
        )

    image_ratio = (
        image_width
        / image_height
    )

    target_ratio = (
        target_width
        / target_height
    )

    picture = slide.shapes.add_picture(
        str(resolved),
        PPTXInches(x),
        PPTXInches(y),
        width=PPTXInches(
            target_width
        ),
        height=PPTXInches(
            target_height
        ),
    )

    if image_ratio > target_ratio:

        crop = (
            1
            - (
                target_ratio
                / image_ratio
            )
        ) / 2

        picture.crop_left = crop
        picture.crop_right = crop

    elif image_ratio < target_ratio:

        crop = (
            1
            - (
                image_ratio
                / target_ratio
            )
        ) / 2

        picture.crop_top = crop
        picture.crop_bottom = crop

    return picture


# ---------------------------------------------------------------------------
# PPTX TITLE SLIDE
# ---------------------------------------------------------------------------

def _pptx_add_title_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
    slide_number: int,
) -> Any:

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _pptx_add_background(
        slide
    )

    left_panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PPTXInches(0.75),
        PPTXInches(1.0),
        PPTXInches(7.95),
        PPTXInches(5.25),
    )

    left_panel.fill.solid()

    left_panel.fill.fore_color.rgb = (
        NOVA_PANEL
    )

    left_panel.fill.transparency = 8

    left_panel.line.color.rgb = (
        NOVA_LINE
    )

    left_panel.line.width = PPTXPt(
        1
    )

    rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PPTXInches(0.75),
        PPTXInches(1.0),
        PPTXInches(0.075),
        PPTXInches(5.25),
    )

    rail.fill.solid()

    rail.fill.fore_color.rgb = (
        NOVA_CYAN
    )

    rail.line.fill.background()

    _pptx_add_accent_chip(
        slide,
        "NOVA / PRESENTATION",
        1.05,
        1.35,
        2.35,
    )

    title_box = slide.shapes.add_textbox(
        PPTXInches(1.05),
        PPTXInches(2.0),
        PPTXInches(7.2),
        PPTXInches(1.8),
    )

    tf = title_box.text_frame

    tf.clear()

    tf.word_wrap = True

    tf.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    paragraph = tf.paragraphs[0]

    paragraph.alignment = (
        PP_ALIGN.LEFT
    )

    _pptx_add_text_runs(
        paragraph=paragraph,
        text=title or "NOVA Presentation",
        font_size=31,
        font_color=NOVA_WHITE,
        font_name="Aptos Display",
        bold_default=True,
    )

    subtitle_box = slide.shapes.add_textbox(
        PPTXInches(1.08),
        PPTXInches(4.0),
        PPTXInches(6.9),
        PPTXInches(1.0),
    )

    stf = subtitle_box.text_frame

    stf.clear()

    stf.word_wrap = True

    sp = stf.paragraphs[0]

    _pptx_add_text_runs(
        paragraph=sp,
        text=(
            subtitle
            or "Sovereign Industrial Intelligence"
        ),
        font_size=15,
        font_color=NOVA_MUTED,
        font_name="Aptos",
    )

    orb_outer = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        PPTXInches(9.35),
        PPTXInches(1.35),
        PPTXInches(2.8),
        PPTXInches(2.8),
    )

    orb_outer.fill.solid()

    orb_outer.fill.fore_color.rgb = (
        NOVA_PANEL_LIGHT
    )

    orb_outer.fill.transparency = 20

    orb_outer.line.color.rgb = (
        NOVA_LINE
    )

    orb_outer.line.width = PPTXPt(
        1.2
    )

    orb_inner = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        PPTXInches(9.75),
        PPTXInches(1.75),
        PPTXInches(2.0),
        PPTXInches(2.0),
    )

    orb_inner.fill.solid()

    orb_inner.fill.fore_color.rgb = (
        NOVA_BG
    )

    orb_inner.line.color.rgb = (
        NOVA_CYAN
    )

    orb_inner.line.width = PPTXPt(
        2.0
    )

    core = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        PPTXInches(10.35),
        PPTXInches(2.35),
        PPTXInches(0.8),
        PPTXInches(0.8),
    )

    core.fill.solid()

    core.fill.fore_color.rgb = (
        NOVA_CYAN
    )

    core.line.fill.background()

    tech_box = slide.shapes.add_textbox(
        PPTXInches(9.15),
        PPTXInches(4.55),
        PPTXInches(3.0),
        PPTXInches(1.0),
    )

    tech_tf = tech_box.text_frame

    tech_tf.clear()

    tech_paragraph = tech_tf.paragraphs[0]

    tech_paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    tech_run = tech_paragraph.add_run()

    tech_run.text = (
        "LOCAL • SECURE • AGENTIC • MULTIMODAL"
    )

    tech_run.font.name = "Aptos"
    tech_run.font.size = PPTXPt(
        8
    )

    tech_run.font.bold = True

    tech_run.font.color.rgb = (
        NOVA_CYAN
    )

    _pptx_add_footer(
        slide,
        slide_number,
    )

    return slide


# ---------------------------------------------------------------------------
# PPTX EXECUTIVE SUMMARY
# ---------------------------------------------------------------------------

def _pptx_add_executive_summary_slide(
    prs: Presentation,
    title: str,
    bullets: List[str],
    slide_number: int,
    visual_path: str | None = None,
) -> Any:

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _pptx_add_background(
        slide
    )

    _pptx_add_section_marker(
        slide,
        "EXECUTIVE SUMMARY",
        0.8,
        0.33,
    )

    _pptx_add_title_text(
        slide,
        title,
        size=26,
        width=10.8
        if visual_path
        else 11.8,
    )

    if visual_path:

        cards = bullets[:3]

        card_width = 2.45
        start_x = 0.8
        gap = 0.20
        y = 2.0
        card_height = 2.0

        for index, item in enumerate(
            cards
        ):

            x = (
                start_x
                + index
                * (
                    card_width
                    + gap
                )
            )

            _pptx_add_panel(
                slide,
                x,
                y,
                card_width,
                card_height,
                fill_color=NOVA_PANEL,
            )

            badge = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                PPTXInches(
                    x + 0.18
                ),
                PPTXInches(
                    y + 0.18
                ),
                PPTXInches(
                    0.45
                ),
                PPTXInches(
                    0.45
                ),
            )

            badge.fill.solid()

            badge.fill.fore_color.rgb = (
                NOVA_CYAN
            )

            badge.line.fill.background()

            badge_tf = badge.text_frame

            badge_tf.clear()

            badge_tf.vertical_anchor = (
                MSO_ANCHOR.MIDDLE
            )

            badge_p = badge_tf.paragraphs[0]

            badge_p.alignment = (
                PP_ALIGN.CENTER
            )

            badge_r = badge_p.add_run()

            badge_r.text = str(
                index + 1
            )

            badge_r.font.name = "Aptos"

            badge_r.font.size = PPTXPt(
                9
            )

            badge_r.font.bold = True

            badge_r.font.color.rgb = (
                NOVA_DARK
            )

            text_box = slide.shapes.add_textbox(
                PPTXInches(
                    x + 0.18
                ),
                PPTXInches(
                    y + 0.82
                ),
                PPTXInches(
                    card_width
                    - 0.36
                ),
                PPTXInches(
                    1.0
                ),
            )

            text_tf = text_box.text_frame

            text_tf.clear()

            text_tf.word_wrap = True

            paragraph = text_tf.paragraphs[0]

            _pptx_add_text_runs(
                paragraph,
                item,
                font_size=12,
                font_color=NOVA_WHITE,
                font_name="Aptos",
            )

        _pptx_add_panel(
            slide,
            8.25,
            1.82,
            4.3,
            4.6,
            fill_color=NOVA_PANEL,
        )

        try:

            _pptx_add_picture_cover(
                slide,
                visual_path,
                8.48,
                2.05,
                3.85,
                4.14,
            )

        except Exception:
            pass

    else:

        cards = bullets[:4]

        card_width = 2.72
        start_x = 0.8
        gap = 0.22
        y = 2.0
        h = 2.05

        for index, item in enumerate(
            cards
        ):

            x = (
                start_x
                + index
                * (
                    card_width
                    + gap
                )
            )

            _pptx_add_panel(
                slide,
                x,
                y,
                card_width,
                h,
            )

            number_box = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                PPTXInches(
                    x + 0.18
                ),
                PPTXInches(
                    y + 0.18
                ),
                PPTXInches(
                    0.48
                ),
                PPTXInches(
                    0.48
                ),
            )

            number_box.fill.solid()

            number_box.fill.fore_color.rgb = (
                NOVA_CYAN
            )

            number_box.line.fill.background()

            tf = number_box.text_frame

            tf.clear()

            tf.vertical_anchor = (
                MSO_ANCHOR.MIDDLE
            )

            np = tf.paragraphs[0]

            np.alignment = (
                PP_ALIGN.CENTER
            )

            nr = np.add_run()

            nr.text = str(
                index + 1
            )

            nr.font.name = "Aptos"

            nr.font.size = PPTXPt(
                9
            )

            nr.font.bold = True

            nr.font.color.rgb = (
                NOVA_DARK
            )

            text_box = slide.shapes.add_textbox(
                PPTXInches(
                    x + 0.2
                ),
                PPTXInches(
                    y + 0.86
                ),
                PPTXInches(
                    card_width
                    - 0.4
                ),
                PPTXInches(
                    1.0
                ),
            )

            ttf = text_box.text_frame

            ttf.clear()

            ttf.word_wrap = True

            tp = ttf.paragraphs[0]

            _pptx_add_text_runs(
                tp,
                item,
                font_size=12.5,
                font_color=NOVA_WHITE,
                font_name="Aptos",
            )

    _pptx_add_footer(
        slide,
        slide_number,
    )

    return slide


# ---------------------------------------------------------------------------
# PPTX VISUAL CONTENT SLIDE
# ---------------------------------------------------------------------------

def _pptx_add_visual_content_slide(
    prs: Presentation,
    title: str,
    bullets: List[str],
    image_path: str,
    slide_number: int,
    section_label: str = "VISUAL INTELLIGENCE",
) -> Any:

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    _pptx_add_background(
        slide
    )

    _pptx_add_section_marker(
        slide,
        section_label,
        0.8,
        0.33,
    )

    _pptx_add_title_text(
        slide,
        title,
        x=0.8,
        y=0.58,
        width=11.8,
        height=0.75,
        size=25,
        color=NOVA_WHITE,
    )

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PPTXInches(0.8),
        PPTXInches(1.42),
        PPTXInches(1.1),
        PPTXInches(0.05),
    )

    line.fill.solid()

    line.fill.fore_color.rgb = (
        NOVA_CYAN
    )

    line.line.fill.background()

    _pptx_add_panel(
        slide,
        0.8,
        1.72,
        6.2,
        4.9,
        fill_color=NOVA_PANEL,
    )

    _pptx_add_panel(
        slide,
        7.23,
        1.72,
        5.32,
        4.9,
        fill_color=NOVA_PANEL,
    )

    _pptx_add_accent_chip(
        slide,
        "KEY INSIGHTS",
        1.13,
        2.02,
        1.65,
    )

    cleaned_items = [
        _pptx_clean_text(
            item
        )
        for item in (
            bullets or []
        )
        if _pptx_clean_text(
            item
        )
    ]

    if not cleaned_items:

        cleaned_items = [
            "NOVA generated this presentation section locally."
        ]

    _pptx_add_bullet_block(
        slide,
        cleaned_items[:6],
        1.1,
        2.58,
        5.45,
        3.58,
        font_size=13.8,
    )

    image_frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        PPTXInches(7.52),
        PPTXInches(2.03),
        PPTXInches(4.74),
        PPTXInches(4.12),
    )

    image_frame.fill.solid()

    image_frame.fill.fore_color.rgb = (
        NOVA_DARK
    )

    image_frame.line.color.rgb = (
        NOVA_LINE
    )

    image_frame.line.width = PPTXPt(
        0.8
    )

    try:

        _pptx_add_picture_cover(
            slide,
            image_path,
            7.59,
            2.10,
            4.60,
            3.98,
        )

    except Exception:
        pass

    telemetry = slide.shapes.add_textbox(
        PPTXInches(7.6),
        PPTXInches(6.22),
        PPTXInches(4.5),
        PPTXInches(0.25),
    )

    telemetry_tf = telemetry.text_frame

    telemetry_tf.clear()

    telemetry_p = telemetry_tf.paragraphs[0]

    telemetry_r = telemetry_p.add_run()

    telemetry_r.text = (
        "LOCAL VISUAL ASSET  /  TOPIC MATCHED"
    )

    telemetry_r.font.name = "Aptos"

    telemetry_r.font.size = PPTXPt(
        7
    )

    telemetry_r.font.bold = True

    telemetry_r.font.color.rgb = (
        NOVA_BLUE
    )

    _pptx_add_footer(
        slide,
        slide_number,
    )

    return slide


# ---------------------------------------------------------------------------
# PPTX WRITER
# ---------------------------------------------------------------------------

def pptx_writer_handler(
    file_path: str,
    title: str = "",
    subtitle: str = "",
    slides: List[Dict[str, Any]] = None,
    content: str = "",
    **kwargs: Any,
) -> Dict[str, Any]:
    """
    Generate a professional PowerPoint presentation.

    Fast visual architecture:
    - one title slide with NOVA orb
    - each content slide gets a topic-aware visual
    - real local SD-Turbo image generation by default
    - deterministic local fallback if generation fails
    - visual assets kept outside workspace/output
    """

    if not file_path or not str(file_path).strip():

        raise ValueError(
            "pptx_writer requires a file_path."
        )

    path = _resolve_workspace_file(
        file_path
    )

    if path.suffix.lower() != ".pptx":

        path = path.with_suffix(
            ".pptx"
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):

        raise ValueError(
            "pptx_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():

        raise ValueError(
            "pptx_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs = Presentation()

    prs.slide_width = PPTXInches(
        13.333
    )

    prs.slide_height = PPTXInches(
        7.5
    )

    clean_title = (
        _pptx_clean_text(
            title
        )
        if title
        else path.stem.replace(
            "_",
            " ",
        ).title()
    )

    clean_subtitle = (
        _pptx_clean_text(
            subtitle
        )
        if subtitle
        else "Sovereign Industrial Intelligence"
    )

    normalized_slide_data: List[
        Dict[str, Any]
    ] = []

    if (
        isinstance(
            slides,
            list,
        )
        and slides
    ):

        for slide_data in slides:

            if not isinstance(
                slide_data,
                dict,
            ):
                continue

            slide_title = _pptx_clean_text(
                slide_data.get(
                    "title",
                    "Overview",
                )
            )

            bullets_raw = slide_data.get(
                "bullets",
                [],
            )

            if not isinstance(
                bullets_raw,
                list,
            ):

                bullets_raw = [
                    bullets_raw
                ]

            bullets = [
                _pptx_clean_text(
                    bullet
                )
                for bullet in bullets_raw
                if _pptx_clean_text(
                    bullet
                )
            ]

            supplied_image_path = (
                slide_data.get(
                    "image_path"
                )
            )

            normalized_slide_data.append(
                {
                    "title": (
                        slide_title
                        or "Overview"
                    ),
                    "bullets": bullets,
                    "image_path": supplied_image_path,
                }
            )

    if not normalized_slide_data:

        records = _pptx_parse_content(
            content
        )

        groups = _pptx_group_records(
            records
        )

        for group in groups:

            normalized_slide_data.append(
                {
                    "title": group.get(
                        "title",
                        "Overview",
                    ),
                    "bullets": group.get(
                        "items",
                        [],
                    ),
                    "image_path": None,
                }
            )

    if not normalized_slide_data:

        normalized_slide_data = [
            {
                "title": "Executive Summary",
                "bullets": [
                    "Presentation generated successfully by NOVA."
                ],
                "image_path": None,
            }
        ]

    # ------------------------------------------------------------------
    # TITLE SLIDE
    # ------------------------------------------------------------------

    slide_counter = 1

    _pptx_add_title_slide(
        prs,
        clean_title,
        clean_subtitle,
        slide_counter,
    )

    slide_counter += 1

    # ------------------------------------------------------------------
    # CONTENT SLIDES
    # ------------------------------------------------------------------

    visual_timings = []

    for index, slide_data in enumerate(
        normalized_slide_data
    ):

        slide_title = slide_data.get(
            "title",
            "Overview",
        )

        bullets = slide_data.get(
            "bullets",
            [],
        )

        supplied_image_path = slide_data.get(
            "image_path"
        )

        image_path = None

        # User supplied image gets priority.
        if supplied_image_path:

            try:

                supplied_resolved = (
                    _resolve_workspace_file(
                        str(
                            supplied_image_path
                        )
                    )
                )

                if supplied_resolved.exists():

                    image_path = str(
                        supplied_resolved
                    )

            except Exception:
                image_path = None

        # Otherwise generate a topic-specific image.
        if image_path is None:

            started = datetime.now()

            image_path = (
                _pptx_generate_visual_asset(
                    title=slide_title,
                    bullets=bullets,
                    slide_index=index + 1,
                )
            )

            visual_timings.append(
                (
                    index + 1,
                    (
                        datetime.now()
                        - started
                    ).total_seconds(),
                    _pptx_visual_profile(
                        slide_title,
                        bullets,
                    ),
                )
            )

        if (
            index == 0
            and len(bullets) <= 4
        ):

            _pptx_add_executive_summary_slide(
                prs=prs,
                title=slide_title,
                bullets=bullets,
                slide_number=slide_counter,
                visual_path=image_path,
            )

        else:

            section_label = (
                f"SECTION {index + 1:02d}"
            )

            if (
                supplied_image_path
                and image_path
            ):

                section_label = (
                    "USER SUPPLIED VISUAL"
                )

            _pptx_add_visual_content_slide(
                prs=prs,
                title=slide_title,
                bullets=bullets,
                image_path=image_path,
                slide_number=slide_counter,
                section_label=section_label,
            )

        slide_counter += 1

    # ------------------------------------------------------------------
    # METADATA
    # ------------------------------------------------------------------

    try:

        prs.core_properties.title = (
            clean_title
        )

        prs.core_properties.subject = (
            "Generated locally by NOVA"
        )

        prs.core_properties.author = (
            "NOVA Sovereign Intelligence"
        )

        prs.core_properties.keywords = (
            "NOVA, AI, industrial automation, "
            "agentic AI, sovereign intelligence, "
            "local image generation, SD Turbo"
        )

        prs.core_properties.comments = (
            "Generated inside NOVA's controlled local workspace. "
            "Topic-aware visuals are rendered locally."
        )

    except Exception:
        pass

    # ------------------------------------------------------------------
    # SAVE
    # ------------------------------------------------------------------

    prs.save(
        str(path)
    )

    if not path.exists():

        raise RuntimeError(
            "PPTX presentation was not created."
        )

    file_size = path.stat().st_size

    if file_size <= 0:

        raise RuntimeError(
            "Generated PPTX presentation is empty."
        )

    verified, msg = verify_artifact(
        path,
        "pptx",
    )

    if not verified:

        raise RuntimeError(
            f"PPTX creation failed verification: {msg}"
        )

    total_visual_time = sum(
        timing[1]
        for timing in visual_timings
    )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": ".pptx",
        "size_bytes": file_size,
        "created": True,
        "verification": msg,
        "workspace": "NOVA",
        "tool": "pptx_writer",
        "slide_count": len(
            prs.slides
        ),
        "visual_generation": (
            "local_sd_turbo_topic_matched"
            if NOVA_IMAGE_GENERATION_ENABLED
            else "local_fallback_topic_matched"
        ),
        "image_model": (
            NOVA_IMAGE_MODEL_ID
            if NOVA_IMAGE_GENERATION_ENABLED
            else None
        ),
        "image_steps": (
            NOVA_IMAGE_STEPS
            if NOVA_IMAGE_GENERATION_ENABLED
            else None
        ),
        "image_size": (
            NOVA_IMAGE_SIZE
            if NOVA_IMAGE_GENERATION_ENABLED
            else None
        ),
        "image_output_size": (
            NOVA_IMAGE_OUTPUT_SIZE
            if NOVA_IMAGE_GENERATION_ENABLED
            else None
        ),
        "visual_generation_seconds": round(
            total_visual_time,
            2,
        ),
    }


# ---------------------------------------------------------------------------
# VISUALIZATION WRITER
# ---------------------------------------------------------------------------

import matplotlib

matplotlib.use(
    "Agg"
)

import matplotlib.pyplot as plt


def visualization_writer_handler(
    file_path: str = "",
    title: str = "Data Visualization",
    chart_type: str = "bar",
    data: Dict[str, Any] = None,
    x_label: str = "",
    y_label: str = "",
    **kwargs: Any,
) -> Dict[str, Any]:

    if not file_path or not str(file_path).strip():

        file_path = (
            f"output/chart_"
            f"{int(datetime.now().timestamp())}.png"
        )

    path = _resolve_workspace_file(
        file_path
    )

    if path.suffix.lower() not in {
        ".png",
        ".jpg",
        ".jpeg",
    }:

        path = path.with_suffix(
            ".png"
        )

    normalized_output = str(
        path.relative_to(
            WORKSPACE_ROOT
        )
    ).replace(
        "\\",
        "/",
    )

    if not normalized_output.lower().startswith(
        "output/"
    ):

        raise ValueError(
            "visualization_writer output must be inside workspace/output/."
        )

    if "/input/" in normalized_output.lower():

        raise ValueError(
            "visualization_writer cannot write into workspace/input/."
        )

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    plt.style.use(
        "dark_background"
    )

    fig, ax = plt.subplots(
        figsize=(8, 4.5),
        dpi=150,
    )

    fig.patch.set_facecolor(
        "#121418"
    )

    ax.set_facecolor(
        "#181A20"
    )

    chart_type = (
        chart_type or "bar"
    ).lower()

    labels = []
    values = []

    if data:

        if isinstance(
            data,
            dict,
        ):

            if (
                "labels" in data
                and "values" in data
            ):

                labels = [
                    str(label)
                    for label in data[
                        "labels"
                    ]
                ]

                values = [
                    float(value)
                    for value in data[
                        "values"
                    ]
                ]

            else:

                labels = [
                    str(key)
                    for key in data.keys()
                ]

                values = [
                    float(value)
                    for value in data.values()
                    if isinstance(
                        value,
                        (
                            int,
                            float,
                        ),
                    )
                ]

        elif isinstance(
            data,
            list,
        ):

            labels = [
                f"Item {index + 1}"
                for index in range(
                    len(data)
                )
            ]

            values = [
                float(value)
                for value in data
                if isinstance(
                    value,
                    (
                        int,
                        float,
                    ),
                )
            ]

    if (
        not labels
        or not values
    ):

        labels = [
            "Category A",
            "Category B",
            "Category C",
            "Category D",
        ]

        values = [
            120,
            240,
            180,
            310,
        ]

    cyan_color = "#00F0FF"

    bar_colors = [
        "#00F0FF",
        "#38BDF8",
        "#818CF8",
        "#C084FC",
        "#F472B6",
    ]

    if chart_type in {
        "bar",
        "column",
    }:

        bars = ax.bar(
            labels,
            values,
            color=bar_colors[
                :len(labels)
            ],
            edgecolor="#0284C7",
            linewidth=1,
        )

        for bar in bars:

            height = bar.get_height()

            ax.annotate(
                f"{height:,.1f}",
                xy=(
                    bar.get_x()
                    + bar.get_width()
                    / 2,
                    height,
                ),
                xytext=(
                    0,
                    3,
                ),
                textcoords="offset points",
                ha="center",
                va="bottom",
                color="#E2E8F0",
                fontsize=8,
            )

    elif chart_type in {
        "line",
        "trend",
    }:

        ax.plot(
            labels,
            values,
            marker="o",
            color=cyan_color,
            linewidth=2.5,
            markersize=6,
            markerfacecolor="#FFFFFF",
        )

        ax.fill_between(
            range(
                len(values)
            ),
            values,
            color=cyan_color,
            alpha=0.15,
        )

    elif chart_type in {
        "pie",
        "donut",
    }:

        wedges, texts, autotexts = ax.pie(
            values,
            labels=labels,
            autopct="%1.1f%%",
            colors=bar_colors[
                :len(labels)
            ],
            wedgeprops=dict(
                width=(
                    0.4
                    if chart_type == "donut"
                    else 1.0
                ),
                edgecolor="#121418",
            ),
        )

        for text in texts:

            text.set_color(
                "#E2E8F0"
            )

        for autotext in autotexts:

            autotext.set_color(
                "#FFFFFF"
            )

    elif chart_type == "scatter":

        ax.scatter(
            range(
                len(values)
            ),
            values,
            color=cyan_color,
            s=60,
            alpha=0.8,
        )

    ax.set_title(
        title,
        color="#00F0FF",
        fontsize=12,
        pad=12,
        fontweight="bold",
    )

    if x_label:

        ax.set_xlabel(
            x_label,
            color="#94A3B8",
            fontsize=9,
        )

    if y_label:

        ax.set_ylabel(
            y_label,
            color="#94A3B8",
            fontsize=9,
        )

    ax.tick_params(
        colors="#94A3B8",
        labelsize=8,
    )

    ax.grid(
        True,
        linestyle="--",
        alpha=0.2,
        color="#475569",
    )

    plt.tight_layout()

    plt.savefig(
        str(path),
        facecolor=fig.get_facecolor(),
        edgecolor="none",
    )

    plt.close(
        fig
    )

    verified, msg = verify_artifact(
        path,
        "image",
    )

    if not verified:

        raise RuntimeError(
            f"Chart creation failed verification: {msg}"
        )

    return {
        "file_path": str(
            path.relative_to(
                WORKSPACE_ROOT
            )
        ).replace(
            "\\",
            "/",
        ),
        "file_name": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "created": True,
        "verification": msg,
        "workspace": "NOVA",
        "tool": "visualization_writer",
    }