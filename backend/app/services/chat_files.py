from pathlib import Path
from uuid import uuid4
import csv
import json
import hashlib
from datetime import datetime, timezone

from PIL import Image
import pytesseract
from openpyxl import load_workbook

from app.knowledge.extractor import extract_text


BASE_DIR = Path(__file__).resolve().parent.parent

CHAT_UPLOADS_DIR = (
    BASE_DIR
    / "knowledge"
    / "chat_uploads"
).resolve()


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".docx",
    ".csv",
    ".xlsx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}


IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}


SPREADSHEET_EXTENSIONS = {
    ".csv",
    ".xlsx",
}


MAX_UPLOAD_SIZE = 20 * 1024 * 1024

MAX_SPREADSHEET_ROWS = 500
MAX_SPREADSHEET_COLUMNS = 50


# ---------------------------------------------------------------------------
# TIME / CHECKSUM
# ---------------------------------------------------------------------------

def _utc_now_iso():
    return datetime.now(
        timezone.utc
    ).isoformat()


def _calculate_sha256(
    file_path: Path,
) -> str:
    digest = hashlib.sha256()

    with file_path.open(
        "rb"
    ) as file_handle:
        for chunk in iter(
            lambda: file_handle.read(
                1024 * 1024
            ),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


# ---------------------------------------------------------------------------
# CHAT FILE PATH SECURITY
# ---------------------------------------------------------------------------

def _resolve_chat_file_path(
    stored_name: str,
) -> Path:
    """
    Resolve a stored chat attachment and ensure it remains inside
    NOVA's controlled chat-upload directory.
    """

    if not stored_name or not str(stored_name).strip():
        raise ValueError(
            "stored_name is required."
        )

    raw_path = Path(
        str(stored_name).strip()
    )

    if raw_path.is_absolute():
        resolved = raw_path.resolve()
    else:
        resolved = (
            CHAT_UPLOADS_DIR / raw_path
        ).resolve()

    try:
        resolved.relative_to(
            CHAT_UPLOADS_DIR
        )

    except ValueError as exc:
        raise PermissionError(
            "Access denied: chat attachment must remain inside NOVA's controlled upload directory."
        ) from exc

    return resolved


# ---------------------------------------------------------------------------
# CHAT ATTACHMENT LOOKUP
# ---------------------------------------------------------------------------

def get_chat_file_path(
    file_id: str,
) -> Path:
    """
    Resolve the real local path for a stored chat attachment.

    This reads the attachment metadata first and then validates the stored
    file path against NOVA's controlled upload directory.
    """

    if not file_id or not str(file_id).strip():
        raise ValueError(
            "file_id is required."
        )

    normalized_file_id = str(
        file_id
    ).strip()

    metadata_path = (
        CHAT_UPLOADS_DIR
        / f"{normalized_file_id}.json"
    ).resolve()

    try:
        metadata_path.relative_to(
            CHAT_UPLOADS_DIR
        )

    except ValueError as exc:
        raise PermissionError(
            "Access denied: invalid chat attachment identifier."
        ) from exc

    if not metadata_path.exists():
        raise FileNotFoundError(
            f"Attachment not found: {normalized_file_id}"
        )

    try:
        metadata = json.loads(
            metadata_path.read_text(
                encoding="utf-8"
            )
        )

    except json.JSONDecodeError as exc:
        raise RuntimeError(
            "Attachment metadata is invalid: "
            f"{normalized_file_id}"
        ) from exc

    stored_name = str(
        metadata.get(
            "stored_name",
            "",
        )
    ).strip()

    if not stored_name:
        raise RuntimeError(
            "Attachment metadata does not contain "
            f"a stored filename: {normalized_file_id}"
        )

    file_path = _resolve_chat_file_path(
        stored_name
    )

    if not file_path.exists():
        raise FileNotFoundError(
            "Stored attachment file is missing: "
            f"{normalized_file_id}"
        )

    if not file_path.is_file():
        raise ValueError(
            "Stored attachment path is not a file: "
            f"{normalized_file_id}"
        )

    return file_path


# ---------------------------------------------------------------------------
# IMAGE OCR
# ---------------------------------------------------------------------------

def _extract_image_text(
    file_path: Path,
) -> str:
    """
    Extract visible text from an uploaded image
    using local Tesseract OCR.

    No external service is used.
    """

    try:
        with Image.open(
            file_path
        ) as image:

            image = image.convert(
                "RGB"
            )

            extracted_text = (
                pytesseract.image_to_string(
                    image,
                    config="--psm 6",
                )
            )

    except Exception as exc:
        raise RuntimeError(
            "Failed to process image with local OCR: "
            f"{exc}"
        ) from exc

    return extracted_text


# ---------------------------------------------------------------------------
# SPREADSHEET HELPERS
# ---------------------------------------------------------------------------

def _normalize_spreadsheet_value(
    value,
) -> str:
    """
    Convert a spreadsheet value into readable text.
    """

    if value is None:
        return ""

    if isinstance(
        value,
        bool,
    ):
        return (
            "True"
            if value
            else "False"
        )

    return str(value)


def _normalize_spreadsheet_row(
    row,
) -> list[str]:
    """
    Normalize one spreadsheet row into a bounded list of strings.
    """

    return [
        _normalize_spreadsheet_value(
            value
        )
        for value in list(row)[
            :MAX_SPREADSHEET_COLUMNS
        ]
    ]


def _spreadsheet_to_text(
    file_path: Path,
    extension: str,
) -> str:
    """
    Extract a compact, structured text representation from
    a CSV or XLSX file.
    """

    if extension == ".csv":
        return _csv_to_text(
            file_path
        )

    if extension == ".xlsx":
        return _xlsx_to_text(
            file_path
        )

    raise ValueError(
        "Unsupported spreadsheet extension: "
        f"{extension}"
    )


def _csv_to_text(
    file_path: Path,
) -> str:
    """
    Read a CSV file locally and convert it into structured text.
    """

    rows = []

    try:
        with file_path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as csv_file:

            sample = csv_file.read(
                4096
            )

            csv_file.seek(
                0
            )

            try:
                dialect = (
                    csv.Sniffer().sniff(
                        sample,
                        delimiters=",;\t|",
                    )
                )

            except csv.Error:
                dialect = csv.excel

            reader = csv.reader(
                csv_file,
                dialect,
            )

            for row in reader:
                rows.append(
                    _normalize_spreadsheet_row(
                        row
                    )
                )

                if len(rows) >= (
                    MAX_SPREADSHEET_ROWS + 1
                ):
                    break

    except UnicodeDecodeError as exc:
        raise ValueError(
            "CSV file is not valid UTF-8 text."
        ) from exc

    except csv.Error as exc:
        raise ValueError(
            "CSV file could not be parsed."
        ) from exc

    if not rows:
        return ""

    headers = rows[0]
    data_rows = rows[1:]

    output = [
        "Spreadsheet type: CSV",
        f"Rows analyzed: {len(data_rows)}",
        f"Columns detected: {len(headers)}",
        "",
        "Columns:",
        ", ".join(
            header
            or f"Column {index + 1}"
            for index, header in enumerate(
                headers
            )
        ),
        "",
        "Data:",
    ]

    for row in data_rows:
        values = []

        width = max(
            len(headers),
            len(row),
        )

        for index in range(
            min(
                width,
                MAX_SPREADSHEET_COLUMNS,
            )
        ):
            header = (
                headers[index]
                if index < len(headers)
                else f"Column {index + 1}"
            )

            header = (
                header
                or f"Column {index + 1}"
            )

            value = (
                row[index]
                if index < len(row)
                else ""
            )

            values.append(
                f"{header}={value}"
            )

        output.append(
            " | ".join(
                values
            )
        )

    return "\n".join(
        output
    )


def _xlsx_to_text(
    file_path: Path,
) -> str:
    """
    Read an XLSX workbook locally and convert each worksheet
    into a bounded structured text representation.
    """

    workbook = None

    try:
        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        output = [
            "Spreadsheet type: XLSX",
            (
                "Worksheet count: "
                f"{len(workbook.worksheets)}"
            ),
        ]

        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                rows.append(
                    _normalize_spreadsheet_row(
                        row
                    )
                )

                if len(rows) >= (
                    MAX_SPREADSHEET_ROWS + 1
                ):
                    break

            if not rows:
                output.extend(
                    [
                        "",
                        (
                            "Worksheet: "
                            f"{worksheet.title}"
                        ),
                        "Rows analyzed: 0",
                        "Columns detected: 0",
                    ]
                )

                continue

            headers = rows[0]
            data_rows = rows[1:]

            output.extend(
                [
                    "",
                    (
                        "Worksheet: "
                        f"{worksheet.title}"
                    ),
                    (
                        "Rows analyzed: "
                        f"{len(data_rows)}"
                    ),
                    (
                        "Columns detected: "
                        f"{len(headers)}"
                    ),
                    "",
                    "Columns:",
                    ", ".join(
                        header
                        or f"Column {index + 1}"
                        for index, header in enumerate(
                            headers
                        )
                    ),
                    "",
                    "Data:",
                ]
            )

            for row in data_rows:
                values = []

                width = max(
                    len(headers),
                    len(row),
                )

                for index in range(
                    min(
                        width,
                        MAX_SPREADSHEET_COLUMNS,
                    )
                ):
                    header = (
                        headers[index]
                        if index < len(headers)
                        else f"Column {index + 1}"
                    )

                    header = (
                        header
                        or f"Column {index + 1}"
                    )

                    value = (
                        row[index]
                        if index < len(row)
                        else ""
                    )

                    values.append(
                        f"{header}={value}"
                    )

                output.append(
                    " | ".join(
                        values
                    )
                )

        return "\n".join(
            output
        )

    except Exception as exc:
        raise RuntimeError(
            "Failed to read XLSX file locally: "
            f"{exc}"
        ) from exc

    finally:
        if workbook is not None:
            workbook.close()


# ---------------------------------------------------------------------------
# GENERAL CHAT TEXT EXTRACTION
# ---------------------------------------------------------------------------

def _extract_chat_text(
    file_path: Path,
    extension: str,
) -> str:
    """
    Extract text from supported chat attachments.

    PDF/DOCX/TXT:
        Use NOVA's existing extractor.

    CSV/XLSX:
        Use NOVA's local spreadsheet reader.

    Images:
        Use local Tesseract OCR.
    """

    if extension in IMAGE_EXTENSIONS:
        return _extract_image_text(
            file_path
        )

    if extension in SPREADSHEET_EXTENSIONS:
        return _spreadsheet_to_text(
            file_path=file_path,
            extension=extension,
        )

    return extract_text(
        str(file_path)
    )


# ---------------------------------------------------------------------------
# REGISTRY SYNCHRONIZATION
# ---------------------------------------------------------------------------

def _register_chat_upload(
    file_id: str,
) -> None:
    """
    Register a newly uploaded chat file in NOVA's real file registry.

    The file is intentionally registered without a vault association.
    It therefore remains a chat upload / Recent Upload until the user
    explicitly adds or moves it into a Knowledge Vault.

    The registry manager reads the authoritative chat metadata itself.
    """

    normalized_file_id = str(
        file_id or ""
    ).strip()

    if not normalized_file_id:
        raise ValueError(
            "File ID is required for registry registration."
        )

    try:
        from app.services.vault_manager import (
            register_chat_file,
        )

        register_chat_file(
            file_id=normalized_file_id
        )

    except Exception as exc:
        raise RuntimeError(
            "Chat file was stored, but the NOVA file registry "
            f"could not register it: {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# SAVE CHAT FILE
# ---------------------------------------------------------------------------

def save_chat_file(
    filename: str,
    file_bytes: bytes,
    content_type: str = "",
) -> dict:
    """
    Save a chat attachment into NOVA's local storage.

    The same physical file is also registered in the NOVA file registry
    with vault_id=None. This allows Recent Uploads to reference the real
    file and later add it to a vault without creating a duplicate.

    Supported:
    - PDF
    - TXT
    - DOCX
    - CSV
    - XLSX
    - PNG
    - JPG
    - JPEG
    - WEBP
    """

    if not filename:
        raise ValueError(
            "No filename provided."
        )

    extension = (
        Path(filename)
        .suffix
        .lower()
    )

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. "
            "Use PDF, TXT, DOCX, CSV, XLSX, PNG, JPG, JPEG, or WEBP."
        )

    if not file_bytes:
        raise ValueError(
            "The uploaded file is empty."
        )

    if len(file_bytes) > MAX_UPLOAD_SIZE:
        raise ValueError(
            "File is too large. Maximum size is 20 MB."
        )

    CHAT_UPLOADS_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    file_id = uuid4().hex

    original_name = (
        Path(filename).name
    )

    stored_name = (
        f"{file_id}_{original_name}"
    )

    file_path = (
        CHAT_UPLOADS_DIR
        / stored_name
    ).resolve()

    try:
        file_path.relative_to(
            CHAT_UPLOADS_DIR
        )

    except ValueError as exc:
        raise PermissionError(
            "Access denied: invalid chat attachment path."
        ) from exc

    file_path.write_bytes(
        file_bytes
    )

    try:
        extracted_text = (
            _extract_chat_text(
                file_path=file_path,
                extension=extension,
            )
        )

    except Exception:
        if file_path.exists():
            file_path.unlink()

        raise

    extracted_text = (
        extracted_text or ""
    ).strip()

    if not extracted_text:
        if file_path.exists():
            file_path.unlink()

        if extension in IMAGE_EXTENSIONS:
            raise ValueError(
                "No readable text was found in the image."
            )

        if extension in SPREADSHEET_EXTENSIONS:
            raise ValueError(
                "No readable spreadsheet data was found in the file."
            )

        raise ValueError(
            "No readable text was found in the file."
        )

    if extension in IMAGE_EXTENSIONS:
        file_type = "image"

    elif extension in SPREADSHEET_EXTENSIONS:
        file_type = "spreadsheet"

    else:
        file_type = "document"

    try:
        stored_relative_path = str(
            file_path.relative_to(
                BASE_DIR.resolve()
            )
        ).replace(
            "\\",
            "/",
        )

    except ValueError:
        stored_relative_path = (
            file_path.name
        )

    size_bytes = (
        file_path.stat().st_size
    )

    sha256 = _calculate_sha256(
        file_path
    )

    created_at = _utc_now_iso()

    metadata = {
        "file_id": file_id,
        "filename": original_name,
        "stored_name": stored_name,
        "stored_relative_path": stored_relative_path,
        "content_type": content_type,
        "extension": extension,
        "file_type": file_type,
        "size_bytes": size_bytes,
        "size": size_bytes,
        "characters": len(
            extracted_text
        ),
        "sha256": sha256,
        "created_at": created_at,
        "uploaded_at": created_at,
        "source": "chat_upload",
        "vault_id": None,
        "status": "active",
        "text": extracted_text,
    }

    metadata_path = (
        CHAT_UPLOADS_DIR
        / f"{file_id}.json"
    ).resolve()

    try:
        metadata_path.relative_to(
            CHAT_UPLOADS_DIR
        )

    except ValueError as exc:
        if file_path.exists():
            file_path.unlink()

        raise PermissionError(
            "Access denied: invalid attachment metadata path."
        ) from exc

    try:
        metadata_path.write_text(
            json.dumps(
                metadata,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        _register_chat_upload(
            file_id
        )

    except Exception:
        if metadata_path.exists():
            metadata_path.unlink()

        if file_path.exists():
            file_path.unlink()

        raise

    return {
        "file_id": file_id,
        "filename": original_name,
        "stored_name": stored_name,
        "stored_relative_path": stored_relative_path,
        "content_type": content_type,
        "extension": extension,
        "file_type": file_type,
        "size": size_bytes,
        "size_bytes": size_bytes,
        "characters": len(
            extracted_text
        ),
        "sha256": sha256,
        "created_at": created_at,
        "uploaded_at": created_at,
        "source": "chat_upload",
        "vault_id": None,
        "status": "active",
    }


# ---------------------------------------------------------------------------
# LOAD CHAT FILE
# ---------------------------------------------------------------------------

def load_chat_file(
    file_id: str,
) -> dict:
    """
    Load locally stored chat attachment metadata and validate
    that its backing file still exists.
    """

    if not file_id or not str(file_id).strip():
        raise ValueError(
            "file_id is required."
        )

    normalized_file_id = str(
        file_id
    ).strip()

    metadata_path = (
        CHAT_UPLOADS_DIR
        / f"{normalized_file_id}.json"
    ).resolve()

    try:
        metadata_path.relative_to(
            CHAT_UPLOADS_DIR
        )

    except ValueError as exc:
        raise PermissionError(
            "Access denied: invalid chat attachment identifier."
        ) from exc

    if not metadata_path.exists():
        raise FileNotFoundError(
            f"Attachment not found: {normalized_file_id}"
        )

    try:
        metadata = json.loads(
            metadata_path.read_text(
                encoding="utf-8"
            )
        )

    except json.JSONDecodeError as exc:
        raise RuntimeError(
            "Attachment metadata is invalid: "
            f"{normalized_file_id}"
        ) from exc

    stored_name = str(
        metadata.get(
            "stored_name",
            "",
        )
    ).strip()

    if not stored_name:
        raise RuntimeError(
            "Attachment metadata does not contain "
            f"a stored filename: {normalized_file_id}"
        )

    file_path = _resolve_chat_file_path(
        stored_name
    )

    if not file_path.exists():
        raise FileNotFoundError(
            "Stored attachment file is missing: "
            f"{normalized_file_id}"
        )

    if not file_path.is_file():
        raise ValueError(
            "Stored attachment path is not a file: "
            f"{normalized_file_id}"
        )

    if not metadata.get(
        "stored_relative_path"
    ):
        try:
            metadata["stored_relative_path"] = str(
                file_path.relative_to(
                    BASE_DIR.resolve()
                )
            ).replace(
                "\\",
                "/",
            )

        except ValueError:
            metadata["stored_relative_path"] = (
                file_path.name
            )

    if not metadata.get(
        "size_bytes"
    ):
        metadata["size_bytes"] = (
            file_path.stat().st_size
        )

    if not metadata.get(
        "sha256"
    ):
        metadata["sha256"] = (
            _calculate_sha256(
                file_path
            )
        )

    return metadata